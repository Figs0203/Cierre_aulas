"""
Módulo de Actualización Directa sobre Archivos Oficiales (Biblioteca EAFIT).

Este módulo ejecuta el cierre directamente sobre los libros Excel oficiales con
máxima responsabilidad y mecanismos de seguridad de nivel institucional:

1. Respaldo Automático Previo (Zero-Data-Loss):
   Copia íntegra con marca de tiempo en la subcarpeta '_backups_cierre/' antes de
   tocar cualquier archivo.

2. Regla Estricta de No-Sobreescritura:
   - Celdas con datos preexistentes: Se respetan intactas y no se tocan.
   - Excepción autorizada: Nombres y apellidos de los estudiantes procesados se
     estandarizan a mayúsculas sostenidas (.upper()).
   - Celdas de cierre (módulos, promedio, estado, certificados): Solo se escriben
     si la celda está estrictamente vacía.

3. Separadores de Clase:
   Borde inferior grueso ('medium', color negro) al final de cada clase en
   Sistematización y en Certificados.

4. Archivo de Notas:
   Se mantiene estrictamente como SOLO LECTURA (no se modifica).

5. Compatibilidad con Tablas Dinámicas (PivotTables):
   openpyxl no soporta tablas dinámicas. Al guardar un archivo que las contenía,
   se eliminan para evitar el mensaje de "reparación" de Excel. El archivo original
   siempre tiene su backup íntegro previo, por lo que no hay pérdida de datos.
"""

from __future__ import annotations

import os
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from xml.etree import ElementTree

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

from config.settings import (
    CERTIFICADO_NO,
    CERTIFICADO_SI,
    CONTROL_AULAS_ALL_SHEETS,
    CONTROL_AULAS_IGNORED_SHEET,
    VALOR_NO_APLICA,
)
from core.models import AulaCierreResult, AulaMetadata, ClaseGroup, StudentMatch
from normalization.dates import (
    format_certificate_date,
    format_ciclo_date,
    format_semester_label,
    parse_date,
)
from normalization.text import format_ciclo, format_docente_coin, normalize_document

# Estilos de bordes oficiales.
# Todos los bordes se dibujan en negro (mismo color en los cuatro lados) para que
# sean visibles y consistentes con el formato de los archivos oficiales.
_BORDER_COLOR = "000000"

_BORDER_THIN = Border(
    left=Side(style="thin", color=_BORDER_COLOR),
    right=Side(style="thin", color=_BORDER_COLOR),
    top=Side(style="thin", color=_BORDER_COLOR),
    bottom=Side(style="thin", color=_BORDER_COLOR),
)

# La celda de fin de clase conserva los cuatro bordes: los tres lados con el mismo
# estilo fino que el resto de la tabla y solo el borde inferior en 'medium' (grueso).
_BORDER_CLASS_SEPARATOR = Border(
    left=Side(style="thin", color=_BORDER_COLOR),
    right=Side(style="thin", color=_BORDER_COLOR),
    top=Side(style="thin", color=_BORDER_COLOR),
    bottom=Side(style="medium", color=_BORDER_COLOR),
)

_FONT_BODY = Font(name="Calibri", size=10)

# Fuente oficial de la tabla de Certificados (confirmado: Zurich Cn BT 11).
_FONT_CERTIFICADO = Font(name="Zurich Cn BT", size=11)


@dataclass
class DirectUpdateReport:
    """Informe detallado de la actualización directa en archivos oficiales."""
    timestamp: str
    backups_created: List[Path] = field(default_factory=list)
    sistematizacion_updated_cells: int = 0
    sistematizacion_names_uppercased: int = 0
    sistematizacion_untouched_cells: int = 0
    certificados_rows_appended: int = 0
    control_aulas_classes_marked: int = 0
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    # Verificación de listas de validación por archivo (SheetBlocksReport.resumen()).
    validaciones_preservadas: Dict[str, str] = field(default_factory=dict)

    @property
    def is_successful(self) -> bool:
        return len(self.errors) == 0


def _strip_pivot_tables(wb: openpyxl.Workbook) -> int:
    """Elimina todas las tablas dinámicas (PivotTables) de un workbook en memoria.

    openpyxl no soporta PivotTables y produce XML inválido al guardar archivos que
    las contienen, lo que provoca el mensaje de "reparación" de Excel.
    El archivo original siempre tiene su backup previo, así que no hay pérdida.

    Args:
        wb: El workbook a sanear (se modifica en lugar).

    Returns:
        Número de tablas dinámicas eliminadas.
    """
    count = 0
    for ws in wb.worksheets:
        if hasattr(ws, '_pivots') and ws._pivots:
            count += len(ws._pivots)
            ws._pivots.clear()
    return count


# Detecta referencias a libros externos en fórmulas (ej. "[1]ABBUEI!$G:$G").
_EXTERNAL_REF_RE = re.compile(r"\[\d+\]")


def _strip_external_links(wb: openpyxl.Workbook) -> int:
    """Elimina las referencias a fórmulas externas (externalLinks) de un workbook.

    openpyxl no preserva correctamente la caché de valores de las referencias
    externas (`externalLink*.xml`), lo que al guardar produce un XML inconsistente
    y Excel muestra el aviso de "reparación". Al eliminar las referencias externas,
    dichos archivos dejan de escribirse y el libro abre limpio.

    Las fórmulas que apuntaban a otros libros pueden quedar sin referencia (Excel
    las marca como no disponibles), pero se conservan los valores ya calculados en
    las celdas que sí se escriben. El backup previo garantiza el original íntegro.

    Args:
        wb: El workbook a sanear (se modifica en el lugar).

    Returns:
        Número de referencias externas eliminadas.
    """
    count = 0
    if hasattr(wb, "_external_links") and wb._external_links:
        count = len(wb._external_links)
        wb._external_links.clear()

    # Eliminar nombres definidos que apunten a libros externos (ej. "[1]Hoja!$A:$A")
    # para que no queden referencias colgantes tras quitar los externalLinks.
    try:
        for name, defined in list(wb.defined_names.items()):
            value = getattr(defined, "value", None) or getattr(defined, "attr_text", None) or ""
            if isinstance(value, str) and _EXTERNAL_REF_RE.search(value):
                wb.defined_names.pop(name, None)
                count += 1
    except Exception:
        pass

    return count


def _resolve_external_link_formulas(wb: openpyxl.Workbook, source_path: Path | str) -> int:
    """Sustituye fórmulas que referencian libros externos por su valor en caché.

    Al guardar con openpyxl y eliminar las referencias externas, las fórmulas que
    apuntaban a otros libros (ej. `=+_xlfn.XLOOKUP(AJ79,[1]ABBUEI!$G:$G,...)`)
    quedan con referencias colgantes y Excel las elimina al "reparar" el libro.

    Para evitarlo, se carga el mismo archivo con `data_only=True` (que devuelve el
    último valor calculado por Excel) y se reemplaza cada fórmula con referencia
    externa por dicho valor.

    Args:
        wb: Workbook cargado con `data_only=False` (se modifica en el lugar).
        source_path: Ruta del archivo fuente para leer los valores en caché.

    Returns:
        Número de celdas con fórmula externa sustituidas.
    """
    try:
        wb_values = openpyxl.load_workbook(source_path, data_only=True, read_only=True)
    except Exception:
        return 0

    replaced = 0
    try:
        for ws in wb.worksheets:
            if ws.title not in wb_values.sheetnames:
                continue
            ws_values = wb_values[ws.title]
            for row in ws.iter_rows():
                for cell in row:
                    val = cell.value
                    if isinstance(val, str) and val.startswith("=") and _EXTERNAL_REF_RE.search(val):
                        cached = ws_values.cell(cell.row, cell.column).value
                        cell.value = cached
                        replaced += 1
    finally:
        wb_values.close()

    return replaced


# ---------------------------------------------------------------------------
# PRESERVACIÓN DE LISTAS DE VALIDACIÓN DE DATOS Y EXTENSIONES DE HOJA
# ---------------------------------------------------------------------------
# openpyxl no soporta (y por tanto descarta al guardar):
#   • Las extensiones de hoja (`extLst`), donde Excel 2010+ guarda las listas de
#     validación modernas (`x14:dataValidations`) y el formato condicional moderno.
#   • En algunos escenarios, el bloque estándar `<dataValidations>`.
#   • Nombres definidos del libro (`definedNames`) que son el origen de muchas
#     listas desplegables (p. ej. `formula1="ListaSI"`).
#
# Estrategia (a prueba de pérdidas):
#   1. Se extraen los bloques del respaldo íntegro del archivo original.
#   2. Se COMBINAN con los del archivo guardado por openpyxl (unión sin duplicar),
#      de modo que nunca se pierda una validación presente en cualquiera de los dos.
#   3. Se reinyectan respetando el orden estricto del esquema OOXML y declarando
#      los namespaces necesarios en la raíz del worksheet.
#   4. Se VERIFICA huella por huella (tipo + rango + fórmula) que ninguna validación
#      del original se haya perdido. Si se detecta alguna pérdida se reintenta con
#      una estrategia "dura" (bloque original literal) y, si aun así falla, se
#      informa explícitamente al usuario en lugar de fallar en silencio.

# Bloque <extLst> de una hoja de cálculo (no anidado). Acepta la forma autocerrada.
_SHEET_EXTLST_RE = re.compile(r"<extLst\b[^>]*>.*?</extLst>|<extLst\b[^>]*/>", re.DOTALL)
# Bloque <ext ...>...</ext> individual dentro de un extLst.
_EXT_BLOCK_RE = re.compile(r"<ext\b[^>]*>.*?</ext>", re.DOTALL)
# Atributo uri="..." de un <ext ...>
_EXT_URI_RE = re.compile(r'<ext\b[^>]*\buri="([^"]+)"')
# Declaración de namespace con prefijo: xmlns:prefijo="uri"
_NS_DECL_RE = re.compile(r'\bxmlns:([A-Za-z_][\w.\-]*)\s*=\s*["\']([^"\']+)["\']')
# Bloque <dataValidations>...</dataValidations> (listas de validación estándar).
# Se admite también la forma autocerrada `<dataValidations count="0"/>`.
_DATA_VALIDATIONS_RE = re.compile(
    r"<dataValidations\b[^>]*>.*?</dataValidations>|<dataValidations\b[^>]*/>", re.DOTALL
)
# Elemento <dataValidation> individual (contenido o autocerrado).
# No confunde `x14:dataValidation` porque exige `<dataValidation` tras el '<'.
_DATA_VALIDATION_RE = re.compile(
    r"<dataValidation\b[^>]*/>|<dataValidation\b[^>]*>.*?</dataValidation>", re.DOTALL
)
# Equivalente moderno dentro de `extLst` (`x14:dataValidation`).
_X14_VALIDATION_RE = re.compile(
    r"<x14:dataValidation\b[^>]*/>|<x14:dataValidation\b[^>]*>.*?</x14:dataValidation>",
    re.DOTALL,
)
# Elemento <definedName> individual del libro.
_DEFINED_NAME_RE = re.compile(
    r"<definedName\b[^>]*/>|<definedName\b[^>]*>.*?</definedName>", re.DOTALL
)
# Bloque <definedNames>...</definedNames> del libro.
_DEFINED_NAMES_BLOCK_RE = re.compile(
    r"<definedNames\b[^>]*>.*?</definedNames>|<definedNames\b[^>]*/>", re.DOTALL
)

# Elementos que, según el esquema CT_Worksheet, van DESPUÉS de <dataValidations>.
# IMPORTANTE: no agregar elementos que vayan ANTES (mergeCells, conditionalFormatting,
# autoFilter, etc.); hacerlo produciría un XML fuera de orden de esquema.
_AFTER_DATA_VALIDATIONS = (
    "<hyperlinks", "<printOptions", "<pageMargins", "<pageSetup", "<headerFooter",
    "<rowBreaks", "<colBreaks", "<customProperties", "<cellWatches", "<ignoredErrors",
    "<smartTags", "<drawing", "<legacyDrawing", "<legacyDrawingHF", "<picture",
    "<oleObjects", "<controls", "<webPublishItems", "<tableParts", "<extLst",
)

# Elementos que, según el esquema CT_Workbook, van DESPUÉS de <definedNames>.
_AFTER_DEFINED_NAMES = (
    "<calcPr", "<oleSize", "<customWorkbookViews", "<pivotCaches", "<smartTagPr",
    "<smartTagTypes", "<webPublishing", "<fileRecoveryPr", "<webPublishObjects", "<extLst",
)


def _replace_or_insert_block(sheet_xml: str, regex: "re.Pattern", block: str, after_tags) -> str:
    """Reemplaza un bloque XML si existe, o lo inserta en la posición correcta del esquema.

    Si el bloque no existe, se inserta antes del primero de `after_tags` (elementos
    que van después según el esquema OOXML); si ninguno existe, antes de </worksheet>.
    """
    m = regex.search(sheet_xml)
    if m:
        return sheet_xml[:m.start()] + block + sheet_xml[m.end():]
    pos = len(sheet_xml)
    for tag in after_tags:
        i = sheet_xml.find(tag)
        if i != -1 and i < pos:
            pos = i
    if pos == len(sheet_xml):
        i = sheet_xml.rfind("</worksheet>")
        if i != -1:
            pos = i
    return sheet_xml[:pos] + block + sheet_xml[pos:]


def _collect_ns_decls(xml: str) -> Dict[str, str]:
    """Recolecta las declaraciones {prefijo: uri} de un fragmento XML."""
    return {m.group(1): m.group(2) for m in _NS_DECL_RE.finditer(xml)}


# Prefijos usados en nombres de elementos (<x14:foo>) y atributos (attr:foo="...").
# Se exige una letra tras ':' para no confundir URIs (p. ej. "http://").
_ELEM_PREFIX_RE = re.compile(r"</?([A-Za-z_][\w.\-]*):[A-Za-z_]")
_ATTR_PREFIX_RE = re.compile(r"\s([A-Za-z_][\w.\-]*):[A-Za-z_]")


def _prefixes_used(xml: str) -> set:
    """Devuelve el conjunto de prefijos de namespace usados en un fragmento XML."""
    prefixes = {m.group(1) for m in _ELEM_PREFIX_RE.finditer(xml)}
    prefixes |= {m.group(1) for m in _ATTR_PREFIX_RE.finditer(xml)}
    prefixes.discard("xmlns")
    return prefixes


def _ensure_ns_decls(sheet_xml: str, ns_decls: Dict[str, str]) -> str:
    """Agrega a la etiqueta <worksheet> las declaraciones de namespace faltantes.

    Necesario cuando el `extLst` reinyectado usa prefijos (p. ej. `x14:`, `xm:`)
    declarados en la raíz del worksheet original, que openpyxl no conserva. Sin
    esto, Excel reporta "Prefijo no declarado".
    """
    if not ns_decls:
        return sheet_xml
    m = re.search(r"<worksheet\b[^>]*>", sheet_xml)
    if not m:
        return sheet_xml
    root_tag = m.group(0)
    existing = _collect_ns_decls(root_tag)
    missing = {p: u for p, u in ns_decls.items() if p not in existing}
    if not missing:
        return sheet_xml
    decls = "".join(f' xmlns:{p}="{u}"' for p, u in missing.items())
    new_tag = root_tag[:-1] + decls + ">"
    return sheet_xml.replace(root_tag, new_tag, 1)


def _read_sheet_names_by_path(zf: "zipfile.ZipFile") -> Dict[str, str]:
    """Mapea el nombre de cada hoja a su ruta interna (p. ej. 'Formato' -> 'xl/worksheets/sheet2.xml').

    El mapeo se obtiene de `xl/workbook.xml` (nombre + r:id) y de
    `xl/_rels/workbook.xml.rels` (r:id -> destino).

    Args:
        zf: Archivo .xlsx abierto como zip en modo lectura.

    Returns:
        Diccionario {nombre_hoja: ruta_dentro_del_zip}.
    """
    mapping: Dict[str, str] = {}
    try:
        workbook_xml = zf.read("xl/workbook.xml").decode("utf-8", "ignore")
        rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")
    except KeyError:
        return mapping

    # rId -> Target (ruta relativa a xl/)
    rel_targets: Dict[str, str] = {}
    for m in re.finditer(r'<Relationship\b[^>]*>', rels_xml):
        tag = m.group(0)
        rid_m = re.search(r'Id="([^"]+)"', tag)
        target_m = re.search(r'Target="([^"]+)"', tag)
        if rid_m and target_m:
            target = target_m.group(1).lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            rel_targets[rid_m.group(1)] = target

    # sheet name + r:id
    for m in re.finditer(r'<sheet\b[^>]*>', workbook_xml):
        tag = m.group(0)
        name_m = re.search(r'name="([^"]+)"', tag)
        rid_m = re.search(r'r:id="([^"]+)"', tag)
        if name_m and rid_m and rid_m.group(1) in rel_targets:
            mapping[name_m.group(1)] = rel_targets[rid_m.group(1)]

    return mapping


@dataclass
class SheetBlocksReport:
    """Resultado de la preservación de listas de validación de un archivo oficial.

    Permite informar al usuario, de forma verificable, cuántas listas desplegables
    (validaciones de datos) existían antes y cuántas quedaron después de aplicar el
    cierre, además de cualquier problema detectado durante la restauración.
    """

    archivo: str = ""
    validaciones_original: int = 0
    validaciones_resultado: int = 0
    validaciones_x14_original: int = 0
    validaciones_x14_resultado: int = 0
    hojas_con_validaciones: int = 0
    nombres_definidos_restaurados: int = 0
    problemas: List[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        """True si no se detectó ninguna pérdida de listas de validación."""
        return not self.problemas

    def resumen(self) -> str:
        """Resumen de una línea, apto para la consola del Modo 5."""
        if not self.validaciones_original and not self.validaciones_x14_original:
            return f"✅ {self.archivo}: sin listas de validación de datos que preservar."
        icono = "✅" if self.ok else "⚠️"
        extra = ""
        if self.nombres_definidos_restaurados:
            extra = f" | nombres definidos restaurados: {self.nombres_definidos_restaurados}"
        return (
            f"{icono} {self.archivo}: listas de validación estándar "
            f"{self.validaciones_original}→{self.validaciones_resultado}, "
            f"modernas (x14) {self.validaciones_x14_original}→{self.validaciones_x14_resultado}"
            f"{extra}"
        )


class _MemZip:
    """Adaptador mínimo para reutilizar `_read_sheet_names_by_path` sobre bytes en memoria."""

    def __init__(self, parts: Dict[str, bytes]):
        self._parts = parts

    def read(self, name: str) -> bytes:
        return self._parts[name]


def _read_sheet_names_by_path_dict(members: Dict[str, bytes]) -> Dict[str, str]:
    """Igual que `_read_sheet_names_by_path` pero sobre un dict de partes ya leídas."""
    return _read_sheet_names_by_path(_MemZip(members))


def _normalize_xml_fragment(fragment: str) -> str:
    """Huella estable de un fragmento XML (insensible a espacios y a `xr:uid`)."""
    txt = re.sub(r'\s+xr:uid="[^"]*"', "", fragment)
    return re.sub(r"\s+", " ", txt).strip()


def _validation_fingerprints(sheet_xml: str) -> set:
    """Conjunto de validaciones de datos presentes en el XML de una hoja.

    Incluye tanto las estándar (`<dataValidation>`) como las modernas
    (`<x14:dataValidation>` dentro de `extLst`). Cada huella conserva el rango
    (`sqref`), el tipo y las fórmulas, de modo que dos listas distintas nunca se
    confunden entre sí.
    """
    items = _DATA_VALIDATION_RE.findall(sheet_xml) + _X14_VALIDATION_RE.findall(sheet_xml)
    return {_normalize_xml_fragment(item) for item in items}


def _is_well_formed(xml: str) -> bool:
    """True si el XML se puede parsear (sin prefijos sin declarar, sin errores)."""
    try:
        ElementTree.fromstring(xml)
        return True
    except Exception:  # noqa: BLE001 - cualquier error de parseo invalida el parche
        return False


def _build_data_validations_block(
    source_block: Optional[str],
    saved_xml: str,
    literal: bool = False,
) -> Optional[str]:
    """Construye el bloque `<dataValidations>` definitivo para una hoja.

    Combina (unión sin duplicar) las validaciones del original con las que openpyxl
    haya escrito en el archivo guardado, de modo que ninguna se pierda.

    Args:
        source_block: Bloque `<dataValidations>` del archivo original (respaldo).
        saved_xml: XML de la hoja tal como la guardó openpyxl.
        literal: Si es True devuelve el bloque original tal cual (estrategia "dura").

    Returns:
        El bloque XML a escribir, o None si el original no tenía validaciones.
    """
    source_items = _DATA_VALIDATION_RE.findall(source_block or "")
    if not source_items:
        return None
    if literal:
        return source_block.strip()

    merged = list(source_items)
    seen = {_normalize_xml_fragment(item) for item in merged}
    saved_m = _DATA_VALIDATIONS_RE.search(saved_xml)
    if saved_m:
        for item in _DATA_VALIDATION_RE.findall(saved_m.group(0)):
            fingerprint = _normalize_xml_fragment(item)
            if fingerprint not in seen:
                seen.add(fingerprint)
                merged.append(item)

    return '<dataValidations count="%d">%s</dataValidations>' % (len(merged), "".join(merged))


def _apply_preserved_blocks(
    sheet_xml: str,
    dv_block: Optional[str],
    ext_block: Optional[str],
    ns_pool: Dict[str, str],
    literal: bool = False,
) -> Tuple[str, List[str]]:
    """Reinyecta en el XML de una hoja los bloques de validación y extensiones.

    Args:
        sheet_xml: XML de la hoja guardada por openpyxl.
        dv_block: Bloque `<dataValidations>` del original (o None).
        ext_block: Bloque `<extLst>` del original (o None).
        ns_pool: Declaraciones `{prefijo: uri}` disponibles para resolver prefijos.
        literal: Estrategia "dura": escribe el bloque de validaciones original tal cual.

    Returns:
        Tupla (XML resultante, lista de problemas detectados).
    """
    problemas: List[str] = []
    new_xml = sheet_xml

    if dv_block:
        block = _build_data_validations_block(dv_block, new_xml, literal=literal)
        if block:
            new_xml = _replace_or_insert_block(
                new_xml, _DATA_VALIDATIONS_RE, block, _AFTER_DATA_VALIDATIONS
            )

    if ext_block:
        unknown = _prefixes_used(ext_block) - set(ns_pool)
        if unknown:
            problemas.append(
                "no se pudieron declarar los prefijos de namespace "
                + ", ".join(sorted(unknown))
            )
        else:
            candidate = _merge_extlst_into_sheet_xml(new_xml, ext_block)
            if candidate is None:
                problemas.append("no se pudo reinyectar el bloque extLst")
            else:
                new_xml = candidate

    if new_xml != sheet_xml:
        new_xml = _ensure_ns_decls(new_xml, ns_pool)

    return new_xml, problemas


def _defined_name_key(item: str) -> Tuple[str, str]:
    """Clave de identidad de un `<definedName>`: (nombre en mayúsculas, hoja local)."""
    name_m = re.search(r'\bname="([^"]*)"', item)
    sheet_m = re.search(r'\blocalSheetId="([^"]*)"', item)
    return (
        (name_m.group(1) if name_m else "").upper(),
        sheet_m.group(1) if sheet_m else "",
    )


def _insert_workbook_block(workbook_xml: str, block: str, after_tags) -> str:
    """Inserta un bloque en `xl/workbook.xml` respetando el orden del esquema OOXML."""
    pos = len(workbook_xml)
    for tag in after_tags:
        i = workbook_xml.find(tag)
        if i != -1 and i < pos:
            pos = i
    if pos == len(workbook_xml):
        i = workbook_xml.rfind("</workbook>")
        if i != -1:
            pos = i
    return workbook_xml[:pos] + block + workbook_xml[pos:]


def _restore_defined_names(
    output_workbook_xml: Optional[str],
    source_workbook_xml: Optional[str],
) -> Tuple[Optional[str], int, List[str]]:
    """Reincorpora al libro guardado los nombres definidos del original que falten.

    Muchas listas desplegables usan un nombre definido como origen
    (`<formula1>ListaSI</formula1>`). Si el nombre desaparece del libro, Excel deja
    de mostrar la lista, por lo que se restauran los que falten sin eliminar ni
    duplicar ninguno y sin reintroducir referencias a libros externos.

    Returns:
        Tupla (XML resultante, cantidad restaurada, lista de problemas).
    """
    if not output_workbook_xml or not source_workbook_xml:
        return output_workbook_xml, 0, []

    problemas: List[str] = []
    try:
        source_block_m = _DEFINED_NAMES_BLOCK_RE.search(source_workbook_xml)
        if not source_block_m:
            return output_workbook_xml, 0, []
        source_items = _DEFINED_NAME_RE.findall(source_block_m.group(0))
        if not source_items:
            return output_workbook_xml, 0, []

        target_block_m = _DEFINED_NAMES_BLOCK_RE.search(output_workbook_xml)
        target_block = target_block_m.group(0) if target_block_m else ""
        target_items = _DEFINED_NAME_RE.findall(target_block)

        existing = {_defined_name_key(item) for item in target_items}
        restored: List[str] = []
        for item in source_items:
            if _EXTERNAL_REF_RE.search(item):
                # Nunca se reintroducen nombres apuntando a libros externos:
                # quedarían colgantes y Excel pediría "reparar" el libro.
                continue
            key = _defined_name_key(item)
            if not key[0] or key in existing:
                continue
            existing.add(key)
            restored.append(item)

        if not restored:
            return output_workbook_xml, 0, problemas

        if target_block and not target_block.endswith("/>"):
            merged_block = target_block.replace(
                "</definedNames>", "".join(restored) + "</definedNames>", 1
            )
            return output_workbook_xml.replace(target_block, merged_block, 1), len(restored), problemas

        merged_block = "<definedNames>%s</definedNames>" % "".join(target_items + restored)
        if target_block:
            return output_workbook_xml.replace(target_block, merged_block, 1), len(restored), problemas

        return _insert_workbook_block(output_workbook_xml, merged_block, _AFTER_DEFINED_NAMES), len(restored), problemas
    except Exception as exc:  # noqa: BLE001 - nunca debe romper el cierre
        problemas.append(f"no se pudieron restaurar los nombres definidos: {exc}")
        return output_workbook_xml, 0, problemas


def _read_zip_parts(path: Path) -> Tuple[Dict[str, bytes], Dict[str, "zipfile.ZipInfo"]]:
    """Lee todas las partes de un .xlsx (contenido + metadatos de cada entrada)."""
    with zipfile.ZipFile(path, "r") as zf:
        members = {name: zf.read(name) for name in zf.namelist()}
        infos = {info.filename: info for info in zf.infolist()}
    return members, infos


def _write_zip_parts(
    path: Path,
    members: Dict[str, bytes],
    infos: Dict[str, "zipfile.ZipInfo"],
) -> None:
    """Reescribe un .xlsx preservando el orden y los metadatos de sus partes."""
    import os as _os
    import tempfile as _tempfile

    fd, tmp_name = _tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    _os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zf_new:
            for name, data in members.items():
                info = infos.get(name)
                if info is not None:
                    zf_new.writestr(info, data)
                else:
                    zf_new.writestr(name, data)
        _os.replace(tmp_path, path)
    finally:
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass


def preserve_sheet_data_blocks(
    output_path: Path | str,
    source_path: Path | str,
) -> SheetBlocksReport:
    """Restaura y VERIFICA las listas de validación de datos del archivo original.

    Debe invocarse inmediatamente después de `wb.save(...)`, usando como
    `source_path` la copia de seguridad íntegra creada antes de la modificación.

    Restaura, hoja por hoja:
      • El bloque `<dataValidations>` (listas de validación estándar).
      • El bloque `<extLst>` (validaciones modernas `x14:dataValidations`).
    Y a nivel de libro:
      • Los `<definedNames>` que sirven de origen a algunas listas.

    Finalmente compara huella por huella (rango + tipo + fórmula) que ninguna
    validación del original se haya perdido. Si detecta pérdidas reintenta con una
    estrategia "dura" y, si aun así falla, lo reporta en `problemas`.

    Args:
        output_path: Ruta del archivo recién guardado por openpyxl (se reescribe).
        source_path: Ruta del original íntegro (respaldo) con las validaciones.

    Returns:
        SheetBlocksReport con los conteos antes/después y los problemas detectados.
    """
    out_path = Path(output_path)
    src_path = Path(source_path)
    report = SheetBlocksReport(archivo=out_path.name)

    if not out_path.exists() or not src_path.exists():
        report.problemas.append(
            "No se pudo verificar la preservación de listas de validación "
            "(falta el archivo guardado o su respaldo)."
        )
        return report

    # ---------------------------------------------------------
    # 1. Extraer bloques y huellas del original (respaldo íntegro)
    # ---------------------------------------------------------
    try:
        with zipfile.ZipFile(src_path, "r") as zf_src:
            source_workbook_xml = zf_src.read("xl/workbook.xml").decode("utf-8", "ignore")
            source_map = _read_sheet_names_by_path(zf_src)
            # Declaraciones de namespace del libro (algunas extensiones usan prefijos
            # declarados sólo en la raíz de xl/workbook.xml, p. ej. `xr`).
            workbook_ns = _collect_ns_decls(source_workbook_xml)
            blocks: Dict[str, Dict[str, Any]] = {}
            for sheet_name, sheet_path in source_map.items():
                try:
                    src_xml = zf_src.read(sheet_path).decode("utf-8", "ignore")
                except KeyError:
                    continue
                dv_m = _DATA_VALIDATIONS_RE.search(src_xml)
                ext_m = _SHEET_EXTLST_RE.search(src_xml)
                if not dv_m and not ext_m:
                    continue
                ns_pool = dict(workbook_ns)
                ns_pool.update(_collect_ns_decls(src_xml))
                blocks[sheet_name] = {
                    "dv": dv_m.group(0) if dv_m else None,
                    "ext": ext_m.group(0) if ext_m else None,
                    "ns": ns_pool,
                    "fingerprints": _validation_fingerprints(src_xml),
                }
    except Exception as exc:  # noqa: BLE001
        report.problemas.append(f"No se pudo leer el respaldo '{src_path.name}': {exc}")
        return report

    report.validaciones_original = sum(
        len(_DATA_VALIDATION_RE.findall(blk["dv"] or "")) for blk in blocks.values()
    )
    report.validaciones_x14_original = sum(
        len(_X14_VALIDATION_RE.findall(blk["ext"] or "")) for blk in blocks.values()
    )
    report.hojas_con_validaciones = sum(1 for blk in blocks.values() if blk["fingerprints"])

    if not blocks:
        return report

    # ---------------------------------------------------------
    # 2. Aplicar la restauración sobre el archivo guardado
    # ---------------------------------------------------------
    try:
        out_members, out_infos = _read_zip_parts(out_path)
    except Exception as exc:  # noqa: BLE001
        report.problemas.append(f"No se pudo leer el archivo guardado '{out_path.name}': {exc}")
        return report

    out_map = _read_sheet_names_by_path_dict(out_members)

    def _build_candidate(literal: bool) -> Tuple[Dict[str, bytes], List[str], int]:
        """Genera una versión candidata del archivo con los bloques reinyectados."""
        candidate = dict(out_members)
        problemas: List[str] = []
        restored_names = 0

        for sheet_name, blk in blocks.items():
            sheet_path = out_map.get(sheet_name)
            if not sheet_path or sheet_path not in candidate:
                problemas.append(
                    f"'{sheet_name}': la hoja no existe en el archivo guardado; "
                    "no se pudieron restaurar sus listas de validación"
                )
                continue
            current = candidate[sheet_path].decode("utf-8", "ignore")
            new_xml, sheet_problems = _apply_preserved_blocks(
                current, blk["dv"], blk["ext"], blk["ns"], literal=literal
            )
            problemas.extend(f"'{sheet_name}': {p}" for p in sheet_problems)
            if new_xml != current:
                if _is_well_formed(new_xml):
                    candidate[sheet_path] = new_xml.encode("utf-8")
                else:
                    problemas.append(
                        f"'{sheet_name}': el XML resultante no es válido; "
                        "se conserva la versión guardada por openpyxl"
                    )

        wb_key = "xl/workbook.xml"
        if wb_key in candidate:
            current_wb = candidate[wb_key].decode("utf-8", "ignore")
            new_wb, restored_names, wb_problems = _restore_defined_names(
                current_wb, source_workbook_xml
            )
            problemas.extend(wb_problems)
            if new_wb is not None and new_wb != current_wb:
                if _is_well_formed(new_wb):
                    candidate[wb_key] = new_wb.encode("utf-8")
                else:
                    problemas.append(
                        "'xl/workbook.xml': los nombres definidos no se pudieron restaurar "
                        "(XML resultante inválido)"
                    )
                    restored_names = 0

        # Verificación huella por huella: ninguna validación debe faltar.
        for sheet_name, blk in blocks.items():
            sheet_path = out_map.get(sheet_name)
            if not sheet_path or sheet_path not in candidate:
                continue
            result_xml = candidate[sheet_path].decode("utf-8", "ignore")
            missing = blk["fingerprints"] - _validation_fingerprints(result_xml)
            if missing:
                problemas.append(
                    f"'{sheet_name}': {len(missing)} validación(es) de datos "
                    "no quedaron en el archivo final"
                )
        return candidate, problemas, restored_names

    candidate, problemas, nombres_restaurados = _build_candidate(literal=False)
    if problemas:
        # Estrategia "dura": escribir el bloque original literal (sin combinar).
        strict_candidate, strict_problemas, strict_nombres = _build_candidate(literal=True)
        if len(strict_problemas) < len(problemas):
            candidate, problemas, nombres_restaurados = (
                strict_candidate,
                strict_problemas,
                strict_nombres,
            )

    report.nombres_definidos_restaurados = nombres_restaurados
    # Se eliminan mensajes repetidos conservando el orden original.
    report.problemas.extend(dict.fromkeys(problemas))

    # ---------------------------------------------------------
    # 3. Escribir el resultado (solo si hubo cambios) y devolver conteos
    # ---------------------------------------------------------
    sheets_involved = {out_map[name] for name in blocks if name in out_map}
    report.validaciones_resultado = sum(
        len(_DATA_VALIDATION_RE.findall(candidate[p].decode("utf-8", "ignore")))
        for p in sheets_involved
        if p in candidate
    )
    report.validaciones_x14_resultado = sum(
        len(_X14_VALIDATION_RE.findall(candidate[p].decode("utf-8", "ignore")))
        for p in sheets_involved
        if p in candidate
    )

    if any(candidate.get(k) != out_members.get(k) for k in set(candidate) | set(out_members)):
        _write_zip_parts(out_path, candidate, out_infos)

    return report


def _restore_sheet_blocks(
    output_path: Path | str,
    source_path: Path | str,
) -> tuple:
    """Compatibilidad histórica: devuelve (validaciones estándar, validaciones x14).

    Mantiene la firma anterior para no romper llamadas existentes; toda la lógica e
    información detallada vive en :func:`preserve_sheet_data_blocks`.
    """
    report = preserve_sheet_data_blocks(output_path, source_path)
    return (report.validaciones_resultado, report.validaciones_x14_resultado)


def _merge_extlst_into_sheet_xml(sheet_xml: str, source_extlst: str) -> Optional[str]:
    """Combina el `extLst` original dentro del XML de una hoja guardada por openpyxl.

    - Si la hoja guardada no tiene `extLst`, se inserta el original antes de `</worksheet>`.
    - Si ya tiene `extLst`, se añaden dentro los `<ext>` del original que no estén
      ya presentes (comparando por `uri`), para evitar duplicados.
    - Se admite la forma autocerrada `<extLst/>` escrita por algunos generadores.

    Returns:
        El XML resultante, o None si no se pudo procesar.
    """
    try:
        existing_m = _SHEET_EXTLST_RE.search(sheet_xml)
        source_exts = _EXT_BLOCK_RE.findall(source_extlst)

        if not existing_m:
            if "</worksheet>" not in sheet_xml:
                return None
            return sheet_xml.replace("</worksheet>", source_extlst + "</worksheet>", 1)

        if not source_exts:
            return sheet_xml

        existing_extlst = existing_m.group(0)
        existing_uris = {u.upper() for u in _EXT_URI_RE.findall(existing_extlst)}

        # Bloques <ext> del origen que aún no están presentes
        missing_exts = [
            block for block in source_exts
            if not any(u.upper() in existing_uris for u in _EXT_URI_RE.findall(block))
        ]
        if not missing_exts:
            return sheet_xml

        if existing_extlst.endswith("/>"):
            # <extLst/> vacío (autocerrado): se abre y se le añaden las extensiones.
            opened = existing_extlst[:-2].rstrip() + ">"
            merged_extlst = opened + "".join(missing_exts) + "</extLst>"
        else:
            merged_extlst = existing_extlst.replace(
                "</extLst>", "".join(missing_exts) + "</extLst>", 1
            )
        return sheet_xml.replace(existing_extlst, merged_extlst, 1)
    except Exception:  # noqa: BLE001
        return None


def create_timestamped_backup(file_path: Path | str, backup_dir: Optional[Path | str] = None) -> Path:
    """Crea una copia de seguridad idéntica del archivo con marca de tiempo.

    Args:
        file_path: Ruta del archivo original a respaldar.
        backup_dir: Carpeta de destino (si es None, crea '_backups_cierre' junto al archivo).

    Returns:
        Ruta al archivo de respaldo creado.
    """
    path = Path(file_path).resolve()
    if not path.exists():
        raise FileNotFoundError(f"No se puede respaldar un archivo inexistente: {path}")

    target_dir = Path(backup_dir).resolve() if backup_dir else path.parent / "_backups_cierre"
    target_dir.mkdir(parents=True, exist_ok=True)

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_name = f"{path.stem}_backup_{timestamp_str}{path.suffix}"
    backup_path = target_dir / backup_name

    shutil.copy2(path, backup_path)
    return backup_path


def find_latest_backup(original_path: Path | str, backup_dir: Optional[Path | str] = None) -> Optional[Path]:
    """Busca el backup más reciente de un archivo oficial.

    Args:
        original_path: Ruta al archivo oficial original.
        backup_dir: Carpeta de backups (si es None, usa '_backups_cierre' junto al archivo).

    Returns:
        Ruta al backup más reciente, o None si no hay ninguno.
    """
    path = Path(original_path).resolve()
    target_dir = Path(backup_dir).resolve() if backup_dir else path.parent / "_backups_cierre"

    if not target_dir.exists():
        return None

    # Buscar archivos que coincidan con el patrón: <stem>_backup_<timestamp><suffix>
    pattern = f"{path.stem}_backup_*{path.suffix}"
    candidates = sorted(target_dir.glob(pattern), reverse=True)  # Más reciente primero (orden lexicográfico del timestamp)
    return candidates[0] if candidates else None


def restore_last_backup(
    file_paths: Dict[str, Path | str],
    backup_dir: Optional[Path | str] = None,
) -> Dict[str, Any]:
    """Restaura el último backup disponible para cada archivo oficial.

    Solo restaura los archivos que tienen un backup disponible (Sistematización,
    Certificados y Control de Aulas). El archivo de Notas nunca se modifica, así
    que no tiene backup y se omite.

    Args:
        file_paths: Diccionario con rutas {'control_aulas', 'sistematizacion', 'certificados', ...}.
        backup_dir: Carpeta donde buscar los backups (por defecto '_backups_cierre').

    Returns:
        Diccionario con los resultados por archivo:
            {'role': {'status': 'ok'|'sin_backup'|'error', 'backup': Path|None, 'mensaje': str}}
    """
    resultados: Dict[str, Any] = {}
    roles_a_restaurar = ["sistematizacion", "certificados", "control_aulas"]

    for role in roles_a_restaurar:
        if role not in file_paths:
            resultados[role] = {"status": "omitido", "backup": None, "mensaje": "Archivo no especificado."}
            continue

        original = Path(file_paths[role]).resolve()
        # Inferir carpeta de backup desde el archivo original si no se especificó
        inferred_backup_dir = backup_dir if backup_dir else original.parent / "_backups_cierre"
        latest = find_latest_backup(original, backup_dir=inferred_backup_dir)

        if latest is None:
            resultados[role] = {
                "status": "sin_backup",
                "backup": None,
                "mensaje": f"No se encontró ningún backup para '{original.name}' en '{inferred_backup_dir}'.",
            }
            continue

        try:
            shutil.copy2(latest, original)
            resultados[role] = {
                "status": "ok",
                "backup": latest,
                "mensaje": f"'{original.name}' restaurado exitosamente desde '{latest.name}'.",
            }
        except Exception as exc:
            resultados[role] = {
                "status": "error",
                "backup": latest,
                "mensaje": f"Error al restaurar '{original.name}': {exc}",
            }

    return resultados


def _preserve_sheet_blocks_and_report(
    report: "DirectUpdateReport",
    etiqueta: str,
    output_path: Path | str,
    source_path: Path | str,
) -> SheetBlocksReport:
    """Restaura y verifica las listas de validación de un archivo ya guardado.

    Nunca lanza excepción: cualquier fallo se registra como advertencia en el
    informe para que el usuario lo vea en consola (y pueda restaurar el respaldo).

    Args:
        report: Informe de la actualización directa (se completa en el lugar).
        etiqueta: Nombre legible del archivo para el informe (ej. 'Sistematización').
        output_path: Archivo recién guardado por openpyxl.
        source_path: Respaldo íntegro creado antes de la modificación.

    Returns:
        El `SheetBlocksReport` con el detalle de la preservación.
    """
    try:
        info = preserve_sheet_data_blocks(output_path, source_path)
    except Exception as exc:  # noqa: BLE001 - la preservación nunca debe romper el cierre
        mensaje = (
            f"{etiqueta}: no se pudo verificar la preservación de las listas de "
            f"validación de datos ({exc})."
        )
        report.warnings.append(mensaje)
        return SheetBlocksReport(archivo=Path(output_path).name, problemas=[mensaje])

    report.validaciones_preservadas[etiqueta] = info.resumen()
    for problema in info.problemas:
        report.warnings.append(f"{etiqueta}: {problema}")
    return info


def apply_direct_cierre(
    file_paths: Dict[str, Path | str],
    result: AulaCierreResult,
    codigo_curso: str,
    metadatos_control: Optional[Dict[str, AulaMetadata]] = None,
    start_consecutivo: int = 1,
    backup_dir: Optional[Path | str] = None,
) -> DirectUpdateReport:
    """Aplica el cierre directamente sobre los archivos oficiales con respaldo previo y seguridad de celdas.

    Args:
        file_paths: Diccionario con rutas {'control_aulas', 'sistematizacion', 'notas', 'certificados'}.
        result: Resultado consolidado del cruce y motor de reglas.
        codigo_curso: Código oficial del curso (ej. 'ABBUEI205').
        metadatos_control: Diccionario opcional {clase_id: AulaMetadata}.
        start_consecutivo: Número correlativo inicial sugerido para Certificados.
        backup_dir: Directorio para las copias de seguridad (opcional).

    Returns:
        DirectUpdateReport con las estadísticas de actualización.
    """
    report = DirectUpdateReport(timestamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    path_sist = Path(file_paths["sistematizacion"]).resolve()
    path_cert = Path(file_paths["certificados"]).resolve()
    path_ctrl = Path(file_paths["control_aulas"]).resolve()

    # ---------------------------------------------------------
    # PASO 1: GENERACIÓN DE COPIAS DE SEGURIDAD (BACKUP)
    # ---------------------------------------------------------
    # Los backups también sirven como fuente de las extensiones de hoja
    # (`extLst`: listas de validación modernas, formato condicional moderno, etc.)
    # que openpyxl descarta al guardar.
    backup_by_role: Dict[str, Path] = {}
    for role, p in [("sistematizacion", path_sist), ("certificados", path_cert), ("control_aulas", path_ctrl)]:
        try:
            bk_path = create_timestamped_backup(p, backup_dir=backup_dir)
            report.backups_created.append(bk_path)
            backup_by_role[role] = bk_path
        except Exception as e:
            report.errors.append(f"Error crítico al crear copia de seguridad de {p.name}: {e}")
            return report

    # ---------------------------------------------------------
    # PASO 2: ACTUALIZACIÓN EN SISTEMATIZACIÓN
    # ---------------------------------------------------------
    try:
        wb_sist = openpyxl.load_workbook(path_sist, data_only=False)
        # Sanear el libro antes de guardar para evitar el aviso de reparación de Excel:
        #  1) sustituir fórmulas con referencias externas por su valor en caché,
        #  2) eliminar tablas dinámicas y referencias externas.
        # El backup previo garantiza la integridad del original.
        _resolve_external_link_formulas(wb_sist, path_sist)
        _strip_pivot_tables(wb_sist)
        _strip_external_links(wb_sist)
        ws_sist = wb_sist["Formato"] if "Formato" in wb_sist.sheetnames else wb_sist.active

        # Recopilar todos los estudiantes ordenados por fila original
        all_matches: List[Tuple[StudentMatch, ClaseGroup]] = []
        for cg in result.clases:
            for m in cg.matches:
                all_matches.append((m, cg))
        all_matches.sort(key=lambda x: x[0].estudiante_sist.row_number)

        # Mapeo de columnas relevantes desde el primer estudiante
        first_es = all_matches[0][0].estudiante_sist if all_matches else None
        col_map = first_es.column_indices if first_es and first_es.column_indices else {}

        col_nombres = col_map.get("nombres")
        col_apellidos = col_map.get("apellidos")
        col_documento = col_map.get("documento")
        max_col_sist = ws_sist.max_column or 55

        for idx, (m, cg) in enumerate(all_matches):
            es = m.estudiante_sist
            r_idx = es.row_number

            # 0. VERIFICACIÓN DE SEGURIDAD DE IDENTIDAD EN FILA DESTINO
            # Comprobar que la fila r_idx corresponda efectivamente a la persona correcta
            row_doc_raw = str(ws_sist.cell(r_idx, col_documento).value or "") if col_documento else ""
            row_doc_norm = normalize_document(row_doc_raw)
            es_doc_norm = normalize_document(es.documento.value if es.documento else "")

            row_nom = str(ws_sist.cell(r_idx, col_nombres).value or "").strip().upper() if col_nombres else ""
            row_ape = str(ws_sist.cell(r_idx, col_apellidos).value or "").strip().upper() if col_apellidos else ""
            es_nom = str(es.nombres.value or "").strip().upper() if es.nombres else ""
            es_ape = str(es.apellidos.value or "").strip().upper() if es.apellidos else ""

            doc_matches = bool(row_doc_norm and es_doc_norm and row_doc_norm == es_doc_norm)
            name_matches = bool(
                (row_nom and es_nom and (row_nom in es_nom or es_nom in row_nom))
                or (row_ape and es_ape and (row_ape in es_ape or es_ape in row_ape))
            )

            if not (doc_matches or name_matches):
                report.errors.append(
                    f"DISCREPANCIA EN FILA {r_idx} DE SISTEMATIZACIÓN: La celda contiene '{row_nom} {row_ape}' "
                    f"(Doc: '{row_doc_raw}'), pero se esperaba calificar a '{es_nom} {es_ape}' (Doc: '{es_doc_norm}'). "
                    f"Fila protegida contra escritura errónea."
                )
                continue

            # 1. Estandarizar nombres y apellidos a MAYÚSCULAS
            if col_nombres:
                cell_nom = ws_sist.cell(r_idx, col_nombres)
                if isinstance(cell_nom.value, str) and cell_nom.value:
                    val_upper = cell_nom.value.strip().upper()
                    if cell_nom.value != val_upper:
                        cell_nom.value = val_upper
                        report.sistematizacion_names_uppercased += 1

            if col_apellidos:
                cell_ape = ws_sist.cell(r_idx, col_apellidos)
                if isinstance(cell_ape.value, str) and cell_ape.value:
                    val_upper = cell_ape.value.strip().upper()
                    if cell_ape.value != val_upper:
                        cell_ape.value = val_upper
                        report.sistematizacion_names_uppercased += 1

            # 2. Celdas de cierre: solo si están estrictamente VACÍAS
            #    Se aceptan alias de claves porque el parser oficial usa
            #    ('estado', 'elaboro', 'codigo_cert', 'envio') mientras que otros
            #    orígenes/datos de prueba usan nombres más explícitos.
            target_values: Dict[int, Any] = {}
            col_keys = [
                ("modulo_1", "modulo_1"),
                ("modulo_2", "modulo_2"),
                ("modulo_3", "modulo_3"),
                ("modulo_4", "modulo_4"),
                ("modulo_5", "modulo_5"),
                ("promedio", "promedio"),
                ("estado", "estado_certificacion"),
                ("elaboro", "elaboro_certificado"),
                ("codigo_cert", "codigo_certificado"),
                ("envio", "envio_certificado"),
                ("fecha_envio", "fecha_envio"),
            ]
            for primary_key, alias_key in col_keys:
                c_idx = col_map.get(primary_key) or col_map.get(alias_key)
                if c_idx and c_idx in es.raw_row_data:
                    target_values[c_idx] = es.raw_row_data[c_idx]

            for c_idx, val_to_write in target_values.items():
                if val_to_write is None:
                    continue
                cell = ws_sist.cell(r_idx, c_idx)
                cur_val = cell.value
                if cur_val is None or str(cur_val).strip() == "":
                    cell.value = val_to_write
                    report.sistematizacion_updated_cells += 1
                else:
                    # Celda con datos preexistentes: NO TOCAR
                    report.sistematizacion_untouched_cells += 1

            # 3. Bordes en TODAS las celdas de la fila procesada.
            #    - Fila normal: caja completa con bordes finos.
            #    - Fila de fin de clase: los mismos tres lados finos + borde inferior grueso.
            current_clase = str(es.clase.value if es.clase and es.clase.value else cg.clase_id).strip()
            is_last_of_class = False
            if idx == len(all_matches) - 1:
                is_last_of_class = True
            else:
                next_m, next_cg = all_matches[idx + 1]
                next_es = next_m.estudiante_sist
                next_clase = str(next_es.clase.value if next_es.clase and next_es.clase.value else next_cg.clase_id).strip()
                if current_clase != next_clase:
                    is_last_of_class = True

            row_border = _BORDER_CLASS_SEPARATOR if is_last_of_class else _BORDER_THIN
            for c in range(1, max_col_sist + 1):
                ws_sist.cell(r_idx, c).border = row_border

        wb_sist.save(path_sist)
        wb_sist.close()
        # Preservar las extensiones de hoja (listas de validación x14, formato
        # condicional moderno, etc.) que openpyxl descarta al guardar.
        _preserve_sheet_blocks_and_report(
            report, "Sistematización", path_sist, backup_by_role.get("sistematizacion", path_sist)
        )
    except Exception as e:
        report.errors.append(f"Error al actualizar Sistematización: {e}")

    # ---------------------------------------------------------
    # PASO 3: ANEXAR EN CERTIFICADOS
    # ---------------------------------------------------------
    try:
        wb_cert = openpyxl.load_workbook(path_cert, data_only=False)
        # Sanear fórmulas externas, tablas dinámicas y referencias externas.
        _resolve_external_link_formulas(wb_cert, path_cert)
        _strip_pivot_tables(wb_cert)
        _strip_external_links(wb_cert)
        if "Códigos" in wb_cert.sheetnames:
            ws_cert = wb_cert["Códigos"]
        elif "Certificados" in wb_cert.sheetnames:
            ws_cert = wb_cert["Certificados"]
        elif "CERTIFICADOS" in wb_cert.sheetnames:
            ws_cert = wb_cert["CERTIFICADOS"]
        else:
            ws_cert = wb_cert.active

        # Encontrar la última fila ocupada con datos reales en columna 1 o 2
        last_row = 1
        max_n = 0
        for r in range(2, (ws_cert.max_row or 1) + 1):
            val_n = ws_cert.cell(r, 1).value
            val_c2 = ws_cert.cell(r, 2).value
            if val_n is not None or val_c2 is not None:
                last_row = max(last_row, r)
                if val_n is not None:
                    try:
                        n_int = int(float(str(val_n).strip()))
                        if n_int > max_n:
                            max_n = n_int
                    except (ValueError, TypeError):
                        pass

        # Determinar consecutivo inicial
        current_n = max(max_n + 1, start_consecutivo)

        # Filtrar solo estudiantes aprobados
        all_approved: List[Tuple[StudentMatch, ClaseGroup]] = []
        for cg in result.clases:
            for m in cg.matches:
                st = str(m.estudiante_sist.estado_calculado or "").strip().lower()
                if "aprob" in st and "no" not in st:
                    all_approved.append((m, cg))
        all_approved.sort(key=lambda item: item[0].estudiante_sist.row_number)

        curr_cert_r = last_row + 1
        for idx, (m, cg) in enumerate(all_approved):
            es = m.estudiante_sist
            meta = (metadatos_control or {}).get(cg.clase_id) or AulaMetadata()
            programa_str = meta.programa_academico.value if meta.programa_academico else ""
            catalogo_str = meta.catalogo.value if meta.catalogo else ""
            f_ini_dt = parse_date(meta.fecha_inicio.value) if meta.fecha_inicio else None
            f_fin_dt = parse_date(meta.fecha_fin.value) if meta.fecha_fin else None

            f_ini_str = format_ciclo_date(f_ini_dt) if f_ini_dt else ""
            f_fin_str = format_ciclo_date(f_fin_dt) if f_fin_dt else ""
            f_envio_cert_str = format_certificate_date(f_fin_dt) if f_fin_dt else ""

            ciclo_val = format_ciclo(cg.clase_id, catalogo_str, programa_str, f_ini_str, f_fin_str).upper()

            formador_lider = ""
            if es.formador_lider and es.formador_lider.value:
                formador_lider = str(es.formador_lider.value).strip()
            elif meta.docente_principal and meta.docente_principal.value:
                formador_lider = str(meta.docente_principal.value).strip()

            docente_coin_val = format_docente_coin(codigo_curso, formador_lider).strip().upper()
            # Columna 'Semestre': periodo académico con formato oficial 'AAAA-S' (ej. '2026-1').
            semestre_val = format_semester_label(f_fin_dt) if f_fin_dt else ""
            ano_val = f_fin_dt.year if isinstance(f_fin_dt, (date, datetime)) else 2026

            nom = str(es.nombres.value if es.nombres else "").strip().upper()
            ape = str(es.apellidos.value if es.apellidos else "").strip().upper()
            doc = es.documento.value if es.documento else ""
            cor = str(es.correo.value if es.correo else "").strip()

            row_cert = [
                current_n,
                ciclo_val,
                docente_coin_val,
                nom,
                ape,
                doc,
                cor,
                str(codigo_curso).strip().upper(),
                f_envio_cert_str,
                ano_val,
                semestre_val,
                1,  # Total certificados = 1
                None,
            ]

            # Separador con borde inferior grueso al terminar cada clase distinta
            current_clase = str(es.clase.value if es.clase and es.clase.value else cg.clase_id).strip()
            is_last_of_class = False
            if idx == len(all_approved) - 1:
                is_last_of_class = True
            else:
                next_m, next_cg = all_approved[idx + 1]
                next_es = next_m.estudiante_sist
                next_clase = str(next_es.clase.value if next_es.clase and next_es.clase.value else next_cg.clase_id).strip()
                if current_clase != next_clase:
                    is_last_of_class = True

            border_cell = _BORDER_CLASS_SEPARATOR if is_last_of_class else _BORDER_THIN

            for c_idx, val in enumerate(row_cert, start=1):
                cell = ws_cert.cell(curr_cert_r, c_idx, val)
                cell.font = _FONT_CERTIFICADO
                # La columna N° (1) es la única que NO lleva borde inferior grueso
                # en la fila de fin de clase; el resto sí.
                if is_last_of_class and c_idx == 1:
                    cell.border = _BORDER_THIN
                else:
                    cell.border = border_cell
                # Toda la información de la tabla de Certificados va centrada.
                cell.alignment = Alignment(horizontal="center", vertical="center")

            current_n += 1
            curr_cert_r += 1
            report.certificados_rows_appended += 1

        wb_cert.save(path_cert)
        wb_cert.close()
        # Preservar las extensiones de hoja que openpyxl descarta al guardar.
        _preserve_sheet_blocks_and_report(
            report, "Certificados", path_cert, backup_by_role.get("certificados", path_cert)
        )
    except Exception as e:
        report.errors.append(f"Error al anexar en Certificados: {e}")

    # ---------------------------------------------------------
    # PASO 4: ACTUALIZAR EN CONTROL DE AULAS
    # ---------------------------------------------------------
    try:
        wb_ctrl = openpyxl.load_workbook(path_ctrl, data_only=False)
        # Sanear fórmulas externas, tablas dinámicas y referencias externas.
        _resolve_external_link_formulas(wb_ctrl, path_ctrl)
        _strip_pivot_tables(wb_ctrl)
        _strip_external_links(wb_ctrl)
        fecha_cierre_hoy = datetime.now().strftime("%d/%m/%Y")

        sheets_to_check = [
            s for s in wb_ctrl.sheetnames
            if s != CONTROL_AULAS_IGNORED_SHEET and (s in CONTROL_AULAS_ALL_SHEETS or s.startswith("Tabla_"))
        ]

        clases_pendientes = {str(cg.clase_id).strip() for cg in result.clases}

        for sheet_name in sheets_to_check:
            ws = wb_ctrl[sheet_name]
            max_col = ws.max_column or 1
            max_row = ws.max_row or 1

            col_clase = None
            col_certificados = None

            for c in range(1, max_col + 1):
                h_val = str(ws.cell(1, c).value or "").strip().upper()
                if h_val == "CLASE":
                    col_clase = c
                elif "CERTIFICADOS" in h_val:
                    col_certificados = c

            if not col_clase or not col_certificados:
                continue

            for r in range(2, max_row + 1):
                c_val = ws.cell(r, col_clase).value
                if c_val is None:
                    continue
                c_str = str(int(c_val)) if isinstance(c_val, float) and c_val.is_integer() else str(c_val).strip()

                if c_str in clases_pendientes:
                    cell_cert = ws.cell(r, col_certificados)
                    if cell_cert.value is None or str(cell_cert.value).strip() == "":
                        cell_cert.value = fecha_cierre_hoy
                        report.control_aulas_classes_marked += 1
                    clases_pendientes.discard(c_str)

        wb_ctrl.save(path_ctrl)
        wb_ctrl.close()
        # Preservar las extensiones de hoja que openpyxl descarta al guardar.
        _preserve_sheet_blocks_and_report(
            report, "Control de Aulas", path_ctrl, backup_by_role.get("control_aulas", path_ctrl)
        )
    except Exception as e:
        report.errors.append(f"Error al actualizar Control de Aulas: {e}")

    return report
