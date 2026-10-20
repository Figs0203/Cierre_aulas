"""
Pruebas unitarias del parser de Notas (Interactiva).

Cubren casos atípicos de estructura por curso:
- Archivos SIN columna 'Sección' (p. ej. cursos CMI).
- Columnas de estado/semáforo con rellenos ámbar/dorado que NO deben marcar a
  los estudiantes como excluidos.
- Reconciliación de la exclusión real (naranja) en columnas de identidad.
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from parsing.notas_parser import parse_notas_archivo


def _build_notas(path: Path, headers, rows, fills=None):
    """Construye un archivo de Notas sintético.

    Args:
        path: Ruta destino.
        headers: Encabezados (fila 1).
        rows: Filas de datos.
        fills: Dict opcional {(row_idx_1based, col_idx_1based): "RRGGBB"}.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Calificaciones"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if fills:
        for (r, c), color in fills.items():
            ws.cell(r, c).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    wb.save(path)
    wb.close()


class TestNotasParserEstructurasAtipicas(unittest.TestCase):

    def test_notas_sin_columna_seccion(self):
        """Sin columna 'Sección', la sección debe quedar en None (no leer una nota)."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Notas_CMI205.xlsx"
            headers = [
                "OrgDefinedId", "Username", "Last Name", "First Name",
                "Calificación Módulo 1", "Calificación Módulo 2", "Calificación Módulo 3",
                "Promedio curso completo", "", "N. Semáforo de secuencia",
            ]
            rows = [
                ["1000000001", "jperez", "Perez Castro", "Juan Camilo", 4.5, 4.0, 4.8, 4.4, None, "Verde"],
                ["1000000002", "mgomez", "Gomez Toro", "Maria Alejandra", 5.0, 4.8, 5.0, 4.9, None, "Verde"],
            ]
            _build_notas(path, headers, rows)

            estudiantes, secciones = parse_notas_archivo(path)
            self.assertEqual(len(estudiantes), 2)
            for es in estudiantes:
                self.assertIsNone(es.seccion, "No debe interpretarse una calificación como sección")
            # Los módulos sí se detectan
            self.assertEqual(set(estudiantes[0].calificaciones_modulos.keys()), {"modulo_1", "modulo_2", "modulo_3"})
            # 'Promedio curso completo' se reconoce como nota final oficial
            self.assertIsNotNone(estudiantes[0].nota_final, "Debe detectar 'Promedio curso completo' como nota final")
            self.assertEqual(estudiantes[0].nota_final.value, 4.4)
            self.assertFalse(any(es.is_orange_fill for es in estudiantes))

    def test_semaforo_ambar_no_marca_excluidos(self):
        """Un relleno ámbar/dorado en la columna de semáforo NO debe excluir al estudiante."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Notas_CMI205.xlsx"
            headers = [
                "OrgDefinedId", "Username", "Last Name", "First Name",
                "Calificación Módulo 1", "Calificación Módulo 2", "Calificación Módulo 3",
                "Promedio curso completo", "", "N. Semáforo de secuencia",
            ]
            rows = [
                [f"100000000{i}", f"user{i}", f"Apellido {i}", f"Nombre {i}", 4.5, 4.0, 4.8, 4.4, None, "Amarillo"]
                for i in range(1, 6)
            ]
            # Relleno ámbar FFC000 en la última columna (semáforo) de TODAS las filas
            fills = {(r, 10): "FFC000" for r in range(2, 7)}
            _build_notas(path, headers, rows, fills=fills)

            estudiantes, _ = parse_notas_archivo(path)
            self.assertEqual(len(estudiantes), 5)
            self.assertEqual(
                sum(1 for es in estudiantes if es.is_orange_fill), 0,
                "El semáforo ámbar no debe marcar al estudiante como excluido",
            )

    def test_naranja_en_identidad_si_marca_excluido(self):
        """El naranja de exclusión en columnas de identidad SÍ debe detectarse."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Notas_CMI205.xlsx"
            headers = [
                "OrgDefinedId", "Username", "Last Name", "First Name",
                "Calificación Módulo 1", "Calificación Módulo 2", "Calificación Módulo 3",
                "Promedio curso completo", "", "N. Semáforo de secuencia",
            ]
            rows = [
                ["1000000001", "user1", "Uno", "Estudiante", 4.5, 4.0, 4.8, 4.4, None, "Amarillo"],
                ["1000000002", "user2", "Dos", "Excluido", 2.0, 1.5, 2.0, 2.0, None, "Rojo"],
            ]
            fills = {
                # Semáforo ámbar para ambos (no debe excluir)
                (2, 10): "FFC000",
                (3, 10): "FFC000",
                # Naranja real sobre la identidad de la fila 3
                (3, 1): "FF8C00", (3, 2): "FF8C00", (3, 3): "FF8C00", (3, 4): "FF8C00",
            }
            _build_notas(path, headers, rows, fills=fills)

            estudiantes, _ = parse_notas_archivo(path)
            excluidos = [es for es in estudiantes if es.is_orange_fill]
            self.assertEqual(len(excluidos), 1)
            self.assertEqual(excluidos[0].first_name.value, "Excluido")


    def test_selecciona_hoja_con_notas_entre_varias(self):
        """Si hay varias hojas, debe elegir la que tiene calificaciones, no la del listado general.

        Caso real CMI: 'Calificaciones' solo tiene identidad + sección (otro curso) y
        'CURSO REALIZADO' tiene módulos y promedio. Debe usar 'CURSO REALIZADO'.
        """
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "CMI120_ECONOMIA_2.xlsx"
            wb = openpyxl.Workbook()
            ws1 = wb.active
            ws1.title = "Calificaciones"
            ws1.append(["OrgDefinedId", "Username", "Last Name", "First Name", "Sección"])
            ws1.append(["9000000001", "otro1", "Apellido", "Nombre", "FI0171 CLASE 4164 E. FINANZAS"])
            ws2 = wb.create_sheet("CURSO REALIZADO")
            ws2.append([
                "OrgDefinedId", "Username", "Last Name", "First Name",
                "Calificación Módulo 1", "Calificación Módulo 2", "Calificación Módulo 3",
                "Promedio curso completo", "", "N.º", "Semáforo de secuencia",
            ])
            ws2.append(["1000000001", "juan.perez1", "Perez Castro", "Juan Camilo", 5.0, 4.5, 4.8, 4.8, None, 1, "Verde"])
            wb.save(path)
            wb.close()

            from parsing.notas_parser import _select_notas_sheet

            wb2 = openpyxl.load_workbook(path)
            self.assertEqual(_select_notas_sheet(wb2), "CURSO REALIZADO")
            wb2.close()

            estudiantes, _ = parse_notas_archivo(path)
            self.assertEqual(len(estudiantes), 1)
            self.assertEqual(estudiantes[0].first_name.value, "Juan Camilo")
            self.assertIsNotNone(estudiantes[0].nota_final)
            self.assertEqual(set(estudiantes[0].calificaciones_modulos.keys()), {"modulo_1", "modulo_2", "modulo_3"})


    def test_detecta_modulo_sin_palabra_modulo(self):
        """Regresión CMI121: 'Cuestionario 1' (sin la palabra 'módulo') debe ser módulo 1.

        Además, la columna de identidad puede llamarse 'Documento' en lugar de
        'OrgDefinedId'.
        """
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "CMI121_DERECHO_2.xlsx"
            headers = [
                "Documento", "Username", "Apellidos", "Nombres", "Sección",
                "Cuestionario 1: estrategia y herramientas de búsqueda",
                "Cuestionario módulo 2: evaluación y organización de la información",
                "Cuestionario módulo 3: uso ético de la información",
                "Nota final", "", "N.º", "Semáforo de secuencia",
            ]
            rows = [["1000000001", "user1", "Apellido", "Nombre", "DR1004-1914-DERECHO",
                     5.0, 3.8, 4.9, 4.4, None, 1, "Verde"]]
            _build_notas(path, headers, rows)

            estudiantes, _ = parse_notas_archivo(path)
            es = estudiantes[0]
            # Identidad por columna 'Documento'
            self.assertEqual(es.org_defined_id.value, "1000000001")
            # Módulo 1 detectado aunque el encabezado no diga 'módulo'
            self.assertEqual(
                {k: v.value for k, v in es.calificaciones_modulos.items()},
                {"modulo_1": 5.0, "modulo_2": 3.8, "modulo_3": 4.9},
            )
            self.assertEqual(es.nota_final.value, 4.4)


if __name__ == "__main__":
    unittest.main()
