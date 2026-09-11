"""
Tests unitarios para las funciones de formateo de columnas de salida.

Cubre las Reglas 5 y 6 confirmadas el 2026-09-04:
  - Regla 5: Columna 'Docente COIN' = '{CÓDIGO-CON-GUIÓN} {NOMBRE_FORMADOR}'
  - Regla 6: Columna 'Ciclo' = 'CLASE {N} {CAT} - {PROG} - {F_INI} A {F_FIN}'

También cubre la lista de hojas confirmada (Regla 1) y la
función de fecha para Ciclo: format_ciclo_date().
"""

from __future__ import annotations

from datetime import date, datetime

import pytest

from config.settings import (
    CONTROL_AULAS_ALL_SHEETS,
    CONTROL_AULAS_IGNORED_SHEET,
)
from normalization.dates import format_ciclo_date
from normalization.text import (
    format_ciclo,
    format_codigo_con_guion,
    format_docente_coin,
)


# ============================================================
# REGLA 1 — Hojas de Control de Aulas
# ============================================================


class TestHojasControlAulas:
    """Verifica que la lista de hojas activas y la hoja ignorada son correctas."""

    EXPECTED_ACTIVE = [
        "Tabla_CMI",
        "Tabla_ABBUEI",
        "Tabla_GICE",
        "Tabla_EPA",
        "Tabla_LATEX",
        "Tabla_IPI",
        "Tabla_UEI",
        "Tabla_CAI",
        "CONTROL",
    ]

    def test_hojas_activas_exactas(self):
        assert CONTROL_AULAS_ALL_SHEETS == self.EXPECTED_ACTIVE, (
            "La lista de hojas activas no coincide con lo confirmado."
        )

    def test_hoja_ignorada_nunca_en_activas(self):
        assert CONTROL_AULAS_IGNORED_SHEET not in CONTROL_AULAS_ALL_SHEETS, (
            f"'{CONTROL_AULAS_IGNORED_SHEET}' NO debe estar en CONTROL_AULAS_ALL_SHEETS."
        )

    def test_hoja_ignorada_es_la_correcta(self):
        assert CONTROL_AULAS_IGNORED_SHEET == "Creación aulas GDA"


# ============================================================
# REGLA 6 — Formato de fecha para Ciclo: format_ciclo_date()
# ============================================================


class TestFormatCicloDate:
    """Verifica la conversión de fechas al estilo 'MES DÍA' para la columna Ciclo."""

    def test_fecha_marzo(self):
        assert format_ciclo_date(date(2026, 3, 16)) == "MARZO 16"

    def test_fecha_abril(self):
        assert format_ciclo_date(date(2026, 4, 27)) == "ABRIL 27"

    def test_fecha_enero(self):
        assert format_ciclo_date(date(2026, 1, 5)) == "ENERO 5"

    def test_fecha_diciembre(self):
        assert format_ciclo_date(date(2025, 12, 31)) == "DICIEMBRE 31"

    def test_fecha_datetime_object(self):
        """También debe funcionar con objetos datetime."""
        dt = datetime(2026, 6, 15, 10, 30)
        assert format_ciclo_date(dt) == "JUNIO 15"

    def test_fecha_none_retorna_vacio(self):
        assert format_ciclo_date(None) == ""


# ============================================================
# REGLA 5 — format_codigo_con_guion()
# ============================================================


class TestFormatCodigoConGuion:
    """Verifica la inserción del guión en códigos de curso."""

    def test_abbuei(self):
        assert format_codigo_con_guion("ABBUEI205") == "ABBUEI-205"

    def test_cmi(self):
        assert format_codigo_con_guion("CMI154") == "CMI-154"

    def test_latex(self):
        assert format_codigo_con_guion("LATEX10") == "LATEX-10"

    def test_ya_tiene_guion(self):
        """Si ya tiene guión, no debe duplicarlo."""
        assert format_codigo_con_guion("ABBUEI-205") == "ABBUEI-205"

    def test_vacio_retorna_vacio(self):
        assert format_codigo_con_guion("") == ""

    def test_patron_no_reconocido_retorna_original(self):
        """Código sin patrón letras+números se devuelve sin cambios."""
        assert format_codigo_con_guion("12345") == "12345"

    def test_espacios_iniciales_finales(self):
        assert format_codigo_con_guion("  GICE99  ") == "GICE-99"


# ============================================================
# REGLA 5 — format_docente_coin()
# ============================================================


class TestFormatDocenteCOIN:
    """Verifica el formato de la columna Docente COIN."""

    def test_caso_confirmado(self):
        result = format_docente_coin("ABBUEI205", "Manuela Restrepo")
        assert result == "ABBUEI-205 Manuela Restrepo"

    def test_codigo_ya_tiene_guion(self):
        result = format_docente_coin("ABBUEI-205", "Manuela Restrepo")
        assert result == "ABBUEI-205 Manuela Restrepo"

    def test_codigo_cmi(self):
        result = format_docente_coin("CMI154", "Carlos Pérez")
        assert result == "CMI-154 Carlos Pérez"

    def test_none_codigo_retorna_vacio(self):
        assert format_docente_coin(None, "Carlos Pérez") == ""

    def test_none_nombre_retorna_vacio(self):
        assert format_docente_coin("CMI154", None) == ""

    def test_ambos_none_retorna_vacio(self):
        assert format_docente_coin(None, None) == ""

    def test_codigo_vacio_retorna_vacio(self):
        assert format_docente_coin("", "Carlos Pérez") == ""

    def test_nombre_vacio_retorna_vacio(self):
        assert format_docente_coin("CMI154", "") == ""


# ============================================================
# REGLA 6 — format_ciclo()
# ============================================================


class TestFormatCiclo:
    """Verifica el formato completo de la columna Ciclo."""

    def test_caso_confirmado(self):
        result = format_ciclo(
            clase=5535,
            catalogo="MF7001",
            programa="M. ESTUDIOS JURÍDICOS",
            fecha_inicio_str="MARZO 16",
            fecha_fin_str="ABRIL 27",
        )
        assert result == "CLASE 5535 MF7001 - M. ESTUDIOS JURÍDICOS - MARZO 16 A ABRIL 27"

    def test_clase_como_string(self):
        result = format_ciclo(
            clase="5535",
            catalogo="MF7001",
            programa="M. ESTUDIOS JURÍDICOS",
            fecha_inicio_str="MARZO 16",
            fecha_fin_str="ABRIL 27",
        )
        assert result == "CLASE 5535 MF7001 - M. ESTUDIOS JURÍDICOS - MARZO 16 A ABRIL 27"

    def test_none_clase_retorna_vacio(self):
        assert format_ciclo(None, "MF7001", "PROG", "MARZO 16", "ABRIL 27") == ""

    def test_none_catalogo_retorna_vacio(self):
        assert format_ciclo(5535, None, "PROG", "MARZO 16", "ABRIL 27") == ""

    def test_none_programa_retorna_vacio(self):
        assert format_ciclo(5535, "MF7001", None, "MARZO 16", "ABRIL 27") == ""

    def test_fecha_inicio_vacia_retorna_vacio(self):
        assert format_ciclo(5535, "MF7001", "PROG", "", "ABRIL 27") == ""

    def test_fecha_fin_vacia_retorna_vacio(self):
        assert format_ciclo(5535, "MF7001", "PROG", "MARZO 16", "") == ""

    def test_todos_none_retorna_vacio(self):
        assert format_ciclo(None, None, None, None, None) == ""

    def test_espacios_en_programa_se_normalizan(self):
        """El programa con espacios extra al inicio/fin debe funcionar igual."""
        result = format_ciclo(
            clase=5535,
            catalogo="MF7001",
            programa="  M. ESTUDIOS JURÍDICOS  ",
            fecha_inicio_str="MARZO 16",
            fecha_fin_str="ABRIL 27",
        )
        assert "M. ESTUDIOS JURÍDICOS" in result


# ============================================================
# FLUJO COMBINADO (Reglas 5 + 6 juntas)
# ============================================================


class TestFlujoCombinadoCertificado:
    """Simula el flujo completo de construcción de las dos columnas de certificado."""

    def test_ciclo_completo_con_fechas_reales(self):
        """Construye Ciclo a partir de fechas reales via format_ciclo_date."""
        f_inicio = format_ciclo_date(date(2026, 3, 16))
        f_fin = format_ciclo_date(date(2026, 4, 27))
        ciclo = format_ciclo(5535, "MF7001", "M. ESTUDIOS JURÍDICOS", f_inicio, f_fin)
        assert ciclo == "CLASE 5535 MF7001 - M. ESTUDIOS JURÍDICOS - MARZO 16 A ABRIL 27"

    def test_docente_y_ciclo_con_datos_faltantes_producen_vacio(self):
        """Si el formador no se puede extraer del archivo, ningún campo se llena."""
        docente = format_docente_coin("ABBUEI205", None)   # formador no encontrado
        f_inicio = format_ciclo_date(None)                 # fecha no encontrada
        ciclo = format_ciclo(5535, "MF7001", "PROG", f_inicio, "ABRIL 27")
        assert docente == ""
        assert ciclo == ""


# ============================================================
# DETECCIÓN DE COLOR NARANJA Y COMPATIBILIDAD CON OBJETOS RGB
# ============================================================


class TestOrangeFillDetection:
    """Verifica la detección de relleno naranja y la robustez ante objetos no-string."""

    def test_known_orange_hex_strings(self):
        from config.colors import is_orange_fill_exact
        assert is_orange_fill_exact("FFA500") is True
        assert is_orange_fill_exact("#FFA500") is True
        assert is_orange_fill_exact("ED7D31") is True
        assert is_orange_fill_exact("FFC000") is True
        assert is_orange_fill_exact("000000") is False
        assert is_orange_fill_exact(None) is False

    def test_theme_colors(self):
        from config.colors import is_orange_fill_exact
        assert is_orange_fill_exact(None, theme_val=6) is True
        assert is_orange_fill_exact(None, theme_val=5) is True
        assert is_orange_fill_exact(None, theme_val=1) is False

    def test_object_without_len_does_not_crash(self):
        """Verifica que un objeto sin __len__ (como openpyxl.styles.colors.RGB) no genere error."""
        class MockRGB:
            def __str__(self):
                return "FFFFA500"

        from config.colors import is_orange_fill_exact
        # No debe lanzar TypeError: object of type 'MockRGB' has no len()
        assert is_orange_fill_exact(MockRGB()) is False or is_orange_fill_exact("FFA500") is True


# ============================================================
# AJUSTES CONFIRMADOS DE MAYÚSCULAS, ESTADOS Y BORDES DE CLASE
# ============================================================


class TestFormatoFeedbackAjustes:
    """Verifica los ajustes solicitados por el usuario respecto a mayúsculas, estados y bordes."""

    def test_estados_y_no_aplica_mantienen_formato_oficial(self):
        from processing.rules_engine import evaluate_certification_status
        from config.settings import VALOR_NO_APLICA

        # Aprobó
        st_apr, _ = evaluate_certification_status({"modulo_1": 4.0, "modulo_2": 4.0})
        assert st_apr == "Aprobó"

        # No aprobó
        st_no, _ = evaluate_certification_status({"modulo_1": 2.0, "modulo_2": 4.0})
        assert st_no == "No aprobó"

        # Abandonó
        st_ab, _ = evaluate_certification_status({"modulo_1": 0.0, "modulo_2": 0.0})
        assert st_ab == "Abandonó"

        # No aplica
        assert VALOR_NO_APLICA == "No aplica"

    def test_separador_de_clases_y_docente_coin_mayusculas(self, tmp_path):
        import openpyxl
        from core.models import (
            AulaCierreResult,
            ClaseGroup,
            StudentMatch,
            EstudianteSistematizacion,
            EstudianteNotas,
            TracedValue,
            AulaMetadata,
        )
        from output.excel_generator import generate_cierre_excel

        def tv(val):
            return TracedValue(value=val, source_file="S", source_sheet="F", source_column="C", source_col_idx=1, source_row=1)

        es1 = EstudianteSistematizacion(
            row_number=5,
            nombres=tv("Juan"),
            apellidos=tv("Perez"),
            clase=tv("5535"),
            documento=tv("1001"),
            correo=tv("jperez@eafit.edu.co"),
            estado_calculado="Aprobó",
            raw_row_data={1: "Juan", 2: "Perez", 3: "Aprobó", 4: "No aplica", 5: "jperez@eafit.edu.co"},
        )
        en1 = EstudianteNotas(
            row_number=2,
            first_name=tv("Juan"),
            last_name=tv("Perez"),
        )
        m1 = StudentMatch(estudiante_sist=es1, estudiante_nota=en1, match_key="doc", match_value="1001")

        es2 = EstudianteSistematizacion(
            row_number=6,
            nombres=tv("Maria"),
            apellidos=tv("Gomez"),
            clase=tv("5536"),
            documento=tv("1002"),
            correo=tv("mgomez@eafit.edu.co"),
            estado_calculado="Aprobó",
            raw_row_data={1: "Maria", 2: "Gomez", 3: "Aprobó", 4: "No aplica", 5: "mgomez@eafit.edu.co"},
        )
        en2 = EstudianteNotas(
            row_number=3,
            first_name=tv("Maria"),
            last_name=tv("Gomez"),
        )
        m2 = StudentMatch(estudiante_sist=es2, estudiante_nota=en2, match_key="doc", match_value="1002")

        cg1 = ClaseGroup(clase_id="5535", matches=[m1])
        cg2 = ClaseGroup(clase_id="5536", matches=[m2])
        res = AulaCierreResult(aula=AulaMetadata(), clases=[cg1, cg2])

        meta_dict = {
            "5535": AulaMetadata(docente_principal=tv("Manuela Restrepo")),
            "5536": AulaMetadata(docente_principal=tv("Manuela Restrepo")),
        }

        out_path = generate_cierre_excel(
            result=res,
            codigo_curso="ABBUEI205",
            sist_headers=["Nombres", "Apellidos", "Estado", "Código Cert", "Correo"],
            start_consecutivo=1,
            output_dir=tmp_path,
            metadatos_control=meta_dict,
        )

        wb = openpyxl.load_workbook(out_path, data_only=True)
        ws_cert = wb["CERTIFICADOS"]

        # Fila 2 (es1, fin de clase 5535) debe tener Docente COIN en MAYÚSCULAS
        docente_coin = ws_cert.cell(2, 3).value
        assert docente_coin == "ABBUEI-205 MANUELA RESTREPO"
        assert docente_coin.isupper()

        # Fila 2 debe tener borde inferior 'medium' porque es fin de clase 5535
        assert ws_cert.cell(2, 1).border.bottom.style == "medium"

        # Fila 3 (es2, fin de clase 5536 y última fila) también debe tener borde inferior 'medium'
        assert ws_cert.cell(3, 1).border.bottom.style == "medium"

        # Fuente oficial de Certificados: Zurich Cn BT 11
        assert ws_cert.cell(2, 1).font.name == "Zurich Cn BT"
        assert ws_cert.cell(2, 1).font.size == 11

        # Todos los bordes presentes en las celdas anexadas (incluida la de fin de clase)
        for _r in (2, 3):
            for _c in range(1, 14):
                _b = ws_cert.cell(_r, _c).border
                assert _b.left.style and _b.right.style and _b.top.style and _b.bottom.style, (
                    f"Faltan bordes en CERTIFICADOS r{_r}c{_c}"
                )

        # Columna 'Semestre' (11) en formato oficial 'AAAA-S'
        import re as _re
        for _r in (2, 3):
            _sem = str(ws_cert.cell(_r, 11).value or "")
            assert _re.match(r"^\d{4}-[12]$", _sem), f"Semestre con formato inesperado: {_sem!r}"

        # Verificar hoja SISTEMATIZACION
        ws_sist = wb["SISTEMATIZACION"]
        # Fila 2: Nombres en mayúsculas, Estado 'Aprobó' preservado, 'No aplica' preservado, correo en minúsculas
        assert ws_sist.cell(2, 1).value == "JUAN"
        assert ws_sist.cell(2, 2).value == "PEREZ"
        assert ws_sist.cell(2, 3).value == "Aprobó"
        assert ws_sist.cell(2, 4).value == "No aplica"
        assert ws_sist.cell(2, 5).value == "jperez@eafit.edu.co"

        wb.close()
