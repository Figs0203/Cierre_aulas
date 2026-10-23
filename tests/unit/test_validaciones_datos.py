"""
Pruebas de preservación de listas de validación de datos (listas desplegables).

Regresión reportada: al aplicar el cierre automático (Modo 5) desaparecían las
listas desplegables de algunas columnas de los archivos oficiales, obligando a
rehacerlas a mano. openpyxl descarta al guardar:

  * Las extensiones de hoja (`extLst`) donde Excel 2010+ guarda las validaciones
    modernas (`x14:dataValidations`).
  * En algunos escenarios, el bloque estándar `<dataValidations>`.
  * Nombres definidos (`definedNames`) que son el origen de varias listas.

Estas pruebas verifican que la herramienta las restaura, que lo hace de forma
idempotente (sin duplicar), que admite las formas autocerradas, que reincorpora
los nombres definidos y que informa el resultado en el reporte del Modo 5.
"""

from __future__ import annotations

import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from core.models import (
    AulaCierreResult,
    AulaMetadata,
    ClaseGroup,
    EstudianteNotas,
    EstudianteSistematizacion,
    StudentMatch,
)
from core.traceability import TracedValue
from output.direct_updater import apply_direct_cierre, preserve_sheet_data_blocks

_SHEET_KEY = "xl/worksheets/sheet1.xml"

_X14_EXT = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}">'
    '<x14:dataValidations count="1">'
    '<x14:dataValidation type="list" allowBlank="1" showErrorMessage="1">'
    '<x14:formula1><xm:f>Listas!$A$1:$A$3</xm:f></x14:formula1>'
    '<xm:sqref>N2:N500</xm:sqref></x14:dataValidation>'
    '</x14:dataValidations></ext></extLst>'
)

_ROOT_NS = (
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" '
    'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"'
)


def _mock_files(tmpdir: Path) -> dict:
    """Reutiliza los archivos sintéticos del suite de actualización directa."""
    # Import local para que pytest no recolecte aquí la clase TestDirectUpdater.
    from tests.unit.test_direct_updater import TestDirectUpdater

    return TestDirectUpdater()._create_mock_files(tmpdir)


def _inject(path: Path, part: str, transform) -> None:
    """Aplica una transformación de texto a una parte interna del .xlsx."""
    with zipfile.ZipFile(path) as z:
        members = {name: z.read(name) for name in z.namelist()}
    members[part] = transform(members[part].decode("utf-8")).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, data in members.items():
            zo.writestr(name, data)


def _read_part(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read(part).decode("utf-8", "ignore")


def _build_result() -> AulaCierreResult:
    """Resultado mínimo (1 estudiante) para ejercitar el cierre directo."""
    col_map = {
        "nombres": 2, "apellidos": 3, "clase": 4, "documento": 5, "correo": 6,
        "modulo_1": 7, "promedio": 12, "estado_certificacion": 13,
        "elaboro_certificado": 14, "codigo_certificado": 15,
        "envio_certificado": 16, "fecha_envio": 17,
    }

    def tv(val, col, row):
        return TracedValue(val, "S", "Formato", col, 1, row)

    es1 = EstudianteSistematizacion(
        row_number=2,
        nombres=tv("carlos andres", "Nombres", 2),
        apellidos=tv("perez gomez", "Apellidos", 2),
        clase=tv("5535", "Clase", 2),
        documento=tv("1001", "Documento", 2),
        correo=tv("cperez@eafit.edu.co", "Correo", 2),
        estado_calculado="Aprobó",
        column_indices=col_map,
        raw_row_data={
            7: 5.0, 12: 4.8, 13: "Aprobó", 14: "SI",
            15: "ABBUEI205", 16: "SI", 17: "27-abr-2026",
        },
    )
    match = StudentMatch(
        estudiante_sist=es1,
        estudiante_nota=EstudianteNotas(row_number=2),
        match_key="doc",
        match_value="1001",
    )
    return AulaCierreResult(
        aula=AulaMetadata(), clases=[ClaseGroup(clase_id="5535", matches=[match])]
    )


def _add_root_ns(xml: str) -> str:
    """Declara los prefijos x14/xm en la raíz del worksheet (como hace Excel)."""
    return xml.replace(
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        + _ROOT_NS,
        1,
    )


def _sheet_with_blocks(path: Path) -> None:
    """Inyecta una validación estándar + una moderna (x14) en la hoja 'Formato'."""
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("N2:N500")
    wb.save(path)
    wb.close()

    def _transform(xml: str) -> str:
        return _add_root_ns(xml).replace("</worksheet>", _X14_EXT + "</worksheet>")

    _inject(path, _SHEET_KEY, _transform)


class TestPreservacionListasDeValidacion(unittest.TestCase):
    """Preservación verificable de listas de validación en el cierre directo."""

    def test_preserva_listas_y_reporta_la_verificacion(self):
        """Tras el cierre las listas siguen presentes y el informe lo confirma."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]

            _sheet_with_blocks(path_sist)

            report = apply_direct_cierre(
                file_paths=files,
                result=_build_result(),
                codigo_curso="ABBUEI205",
                start_consecutivo=1,
            )
            self.assertTrue(report.is_successful, report.errors)
            self.assertEqual(report.warnings, [], report.warnings)

            # El informe expone la verificación de los 3 archivos oficiales
            self.assertIn("Sistematización", report.validaciones_preservadas)
            self.assertIn("Certificados", report.validaciones_preservadas)
            self.assertIn("Control de Aulas", report.validaciones_preservadas)
            resumen = report.validaciones_preservadas["Sistematización"]
            self.assertIn("✅", resumen)
            self.assertIn("estándar 1→1", resumen)
            self.assertIn("modernas (x14) 1→1", resumen)

            content = _read_part(path_sist, _SHEET_KEY)
            self.assertEqual(content.count("<dataValidation "), 1)
            self.assertIn("x14:dataValidations", content)
            self.assertIn("Listas!$A$1:$A$3", content)
            self.assertEqual(content.count("<extLst"), 1)

            with zipfile.ZipFile(path_sist) as z:
                self.assertIsNone(z.testzip())

            # El libro resultante sigue abriendo sin problemas
            wb = openpyxl.load_workbook(path_sist)
            self.assertIn("Formato", wb.sheetnames)
            wb.close()

    def test_escribe_la_fecha_de_envio_como_texto_oficial(self):
        """La columna 'Fecha de envío' oficial queda como 'DD-mes-AAAA' (sin hora)."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]

            report = apply_direct_cierre(
                file_paths=files,
                result=_build_result(),
                codigo_curso="ABBUEI205",
                start_consecutivo=1,
            )
            self.assertTrue(report.is_successful, report.errors)

            wb = openpyxl.load_workbook(path_sist)
            ws = wb["Formato"] if "Formato" in wb.sheetnames else wb.active
            # Columna 17 = 'Fecha de envío' en el archivo sintético
            valor = ws.cell(2, 17).value
            wb.close()

            self.assertEqual(valor, "27-abr-2026")
            self.assertNotIn("00:00", str(valor))
            self.assertNotIn(":", str(valor))

    def test_sqref_multirrango_no_se_confunde_con_prefijo_de_namespace(self):
        """Regresión real (CMI): `sqref="B2:B9 AD2:AF2"` se tomaba como prefijos.

        El detector de prefijos confundía los rangos de celdas de un `sqref`
        multirrango (`AD2:AF2`) con prefijos de namespace, por lo que el `extLst`
        completo se descartaba y se perdían las 8 listas modernas (x14).
        """
        from output.direct_updater import _prefixes_used

        # El fragmento sólo usa dos prefijos reales: x14 y xm
        frag = (
            '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}">'
            '<x14:dataValidations count="1">'
            '<x14:dataValidation type="list">'
            '<x14:formula1><xm:f>Listas!$A$1:$A$3</xm:f></x14:formula1>'
            '<xm:sqref>B2:B9 AD2:AF2 AF2:AG5</xm:sqref>'
            '</x14:dataValidation></x14:dataValidations></ext></extLst>'
        )
        self.assertEqual(_prefixes_used(frag), {"x14", "xm"})

        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            x14_multi = (
                '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}">'
                '<x14:dataValidations count="1">'
                '<x14:dataValidation type="list" allowBlank="1">'
                '<x14:formula1><xm:f>Listas!$A$1:$A$3</xm:f></x14:formula1>'
                '<xm:sqref>B2:B9 AD2:AF2</xm:sqref></x14:dataValidation>'
                '</x14:dataValidations></ext></extLst>'
            )
            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: _add_root_ns(xml).replace("</worksheet>", x14_multi + "</worksheet>"),
            )
            shutil.copy2(path_sist, backup)

            wb = openpyxl.load_workbook(path_sist)
            wb.save(path_sist)  # openpyxl descarta el extLst
            wb.close()

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [], info.problemas)
            self.assertEqual(info.validaciones_x14_original, 1)
            self.assertEqual(info.validaciones_x14_resultado, 1)

            content = _read_part(path_sist, _SHEET_KEY)
            self.assertIn("x14:dataValidations", content)
            self.assertIn("B2:B9 AD2:AF2", content)

    def test_no_duplica_listas_que_openpyxl_reescribe(self):
        """Regresión real: openpyxl normaliza atributos y la unión duplicaba listas.

        openpyxl reescribe `<dataValidation>` (añade showDropDown="0", reordena
        atributos, re-codifica entidades). Al combinar bloques por texto se
        contaban como distintas y las listas se duplicaban (3→6 y 7→14).
        """
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            dv_block = (
                '<dataValidations count="2">'
                '<dataValidation type="list" allowBlank="1" showInputMessage="1" '
                'showErrorMessage="1" errorTitle="Valor no válido" '
                'error="Seleccione de la lista" promptTitle="Seleccione" '
                'prompt="Elija un valor" sqref="N2:N500">'
                '<formula1>&quot;SI,NO&quot;</formula1></dataValidation>'
                '<dataValidation type="list" allowBlank="1" sqref="O2:O500">'
                '<formula1>ListaSI</formula1></dataValidation></dataValidations>'
            )
            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: xml.replace("<pageMargins", dv_block + "<pageMargins", 1),
            )
            shutil.copy2(path_sist, backup)

            # openpyxl reescribe el bloque (atributos normalizados y reordenados)
            wb = openpyxl.load_workbook(path_sist)
            wb.save(path_sist)
            wb.close()
            self.assertEqual(_read_part(path_sist, _SHEET_KEY).count("<dataValidation "), 2)

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [], info.problemas)
            self.assertEqual(info.validaciones_original, 2)
            self.assertEqual(info.validaciones_resultado, 2)

            content = _read_part(path_sist, _SHEET_KEY)
            self.assertEqual(content.count("<dataValidation "), 2)
            self.assertEqual(content.count("<dataValidations"), 1)
            # El bloque debe quedar exactamente como el del respaldo
            self.assertIn('errorTitle="Valor no válido"', content)
            self.assertIn("&quot;SI,NO&quot;", content)

    def test_preservacion_idempotente_sin_duplicar(self):
        """Aplicar la preservación dos veces no duplica ni pierde validaciones."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: _add_root_ns(xml).replace("</worksheet>", _X14_EXT + "</worksheet>"),
            )
            shutil.copy2(path_sist, backup)

            wb = openpyxl.load_workbook(path_sist)
            wb.save(path_sist)  # openpyxl descarta el extLst
            wb.close()

            first = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(first.problemas, [])
            self.assertEqual(first.validaciones_x14_original, 1)
            self.assertEqual(first.validaciones_x14_resultado, 1)

            second = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(second.problemas, [])
            self.assertEqual(second.validaciones_x14_resultado, 1)

            content = _read_part(path_sist, _SHEET_KEY)
            self.assertEqual(content.count("<extLst"), 1)
            self.assertEqual(content.count("<x14:dataValidation "), 1)

    def test_admite_bloques_autocerrados(self):
        """Soporta `<dataValidations/>` y `<extLst/>` (formas autocerradas)."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            dv_block = (
                '<dataValidations count="1">'
                '<dataValidation type="list" allowBlank="1" sqref="M2:M50">'
                '<formula1>"SI,NO,No aplica"</formula1></dataValidation></dataValidations>'
            )
            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: _add_root_ns(xml).replace(
                    "</worksheet>", dv_block + _X14_EXT + "</worksheet>"
                ),
            )
            shutil.copy2(path_sist, backup)

            # Simular la salida de openpyxl: bloques autocerrados y sin contenido
            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: xml.replace(dv_block, '<dataValidations count="0"/>')
                               .replace(_X14_EXT, "<extLst/>"),
            )
            self.assertIn('<dataValidations count="0"/>', _read_part(path_sist, _SHEET_KEY))

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [])

            content = _read_part(path_sist, _SHEET_KEY)
            self.assertIn('sqref="M2:M50"', content)
            self.assertNotIn("<extLst/>", content)
            self.assertIn("x14:dataValidations", content)
            self.assertEqual(content.count("<extLst"), 1)

    def test_restaura_nombres_definidos_de_las_listas(self):
        """Los `definedNames` (origen de listas) se reincorporan si faltan."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            dv_block = (
                '<dataValidations count="1">'
                '<dataValidation type="list" allowBlank="1" sqref="M2:M50">'
                '<formula1>ListaEstados</formula1></dataValidation></dataValidations>'
            )
            _inject(
                path_sist, _SHEET_KEY,
                lambda xml: xml.replace("</worksheet>", dv_block + "</worksheet>"),
            )

            nombre = '<definedName name="ListaEstados">Listas!$A$1:$A$3</definedName>'

            def _add_name(xml: str) -> str:
                if "<definedNames>" in xml:
                    return xml.replace("</definedNames>", nombre + "</definedNames>", 1)
                bloque = "<definedNames>" + nombre + "</definedNames>"
                if "<calcPr" in xml:
                    return xml.replace("<calcPr", bloque + "<calcPr", 1)
                return xml.replace("</workbook>", bloque + "</workbook>", 1)

            _inject(path_sist, "xl/workbook.xml", _add_name)
            shutil.copy2(path_sist, backup)

            # Simular la pérdida del nombre definido al guardar con openpyxl
            def _drop_name(xml: str) -> str:
                return xml.replace(nombre, "")

            _inject(path_sist, "xl/workbook.xml", _drop_name)
            self.assertNotIn("ListaEstados", _read_part(path_sist, "xl/workbook.xml"))

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [])
            self.assertEqual(info.nombres_definidos_restaurados, 1)

            wb_xml = _read_part(path_sist, "xl/workbook.xml")
            self.assertEqual(wb_xml.count('name="ListaEstados"'), 1)

    def test_no_reintroduce_nombres_con_referencias_externas(self):
        """Nunca se restauran nombres que apuntan a libros externos ([N])."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"

            nombre_externo = '<definedName name="Externo">[1]ABBUEI!$G:$G</definedName>'

            def _add_name(xml: str) -> str:
                bloque = "<definedNames>" + nombre_externo + "</definedNames>"
                if "<calcPr" in xml:
                    return xml.replace("<calcPr", bloque + "<calcPr", 1)
                return xml.replace("</workbook>", bloque + "</workbook>", 1)

            _inject(path_sist, "xl/workbook.xml", _add_name)
            shutil.copy2(path_sist, backup)
            _inject(path_sist, "xl/workbook.xml", lambda xml: xml.replace(nombre_externo, ""))

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [])
            self.assertEqual(info.nombres_definidos_restaurados, 0)
            self.assertNotIn("Externo", _read_part(path_sist, "xl/workbook.xml"))

    def test_reporta_problema_si_falta_el_respaldo(self):
        """Sin respaldo no hay verificación posible: se informa, nunca se rompe."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            info = preserve_sheet_data_blocks(
                files["sistematizacion"], tmpdir / "no_existe.xlsx"
            )
            self.assertFalse(info.ok)
            self.assertTrue(info.problemas)
            self.assertIn("No se pudo verificar", info.problemas[0])

    def test_archivo_sin_validaciones_no_reporta_problemas(self):
        """Si el original no tiene listas, no hay nada que preservar ni que fallar."""
        with tempfile.TemporaryDirectory() as tmp_str:
            tmpdir = Path(tmp_str)
            files = _mock_files(tmpdir)
            path_sist = files["sistematizacion"]
            backup = tmpdir / "respaldo.xlsx"
            shutil.copy2(path_sist, backup)

            info = preserve_sheet_data_blocks(path_sist, backup)
            self.assertEqual(info.problemas, [])
            self.assertEqual(info.validaciones_original, 0)
            self.assertEqual(info.validaciones_resultado, 0)


if __name__ == "__main__":
    unittest.main()
