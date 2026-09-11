"""
Generador de Fixtures Sintéticos para Pruebas del Asistente COIN.

Crea 4 libros de Excel sintéticos con datos ficticios para probar:
- Inspección técnica de Fase 0 (hojas, encabezados, fórmulas, colores, formatos).
- Detección de celdas naranjas (estudiantes excluidos).
- Múltiples clases con distintos docentes.
- Fórmulas de estado de certificación.
- Orden aleatorio de estudiantes.
- Aulas ya cerradas y aulas activas.

IMPORTANTE:
    Todos los datos (nombres, documentos, correos) son 100% ficticios
    y no corresponden a personas reales.
"""

from __future__ import annotations

from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


FIXTURES_DIR = Path(__file__).resolve().parent


def create_synthetic_control_aulas(path: Path) -> None:
    wb = openpyxl.Workbook()

    headers = [
        "AÑO", "TIPO DE AULA", "CURSO QUE REALIZA", "PROGRAMA ACADÉMICO",
        "CATÁLOGO", "CLASE", "NOMBRE DE LA SECCIÓN", "DOCENTE PRINCIPAL",
        "FECHA DE INICIO DEL CURSO", "FECHA DE FINALIZACIÓN DEL CURSO",
        "FECHA 2 DE FINALIZACIÓN DEL CURSO", "ENVÍO DE NOTAS A PROFESORES",
        "CERTIFICADOS Y CIERRE DE CURSO", "CERTIFICADOS",
    ]

    sheet_names = [
        "Tabla_CMI",
        "Tabla_ABBUEI",
        "Tabla_GICE",
        "Tabla_EPA",
        "Tabla_LATEX",
        "Tabla_IPI",
        "Tabla_UEI",
        "Tabla_CAI",
        "Creación aulas GDA",
        "CONTROL",
    ]

    # Filas para el aula activa ABBUEI154 y un aula cerrada en Tabla_ABBUEI
    abbuei_rows = [
        [2026, "VIRTUAL", "Búsqueda y uso ético", "Ingeniería", "CAT101", "001", "ABBUEI154-001", "Prof. Ana Gomez", "2026-02-01", "2026-06-16", None, "2026-06-17", "2026-06-18", None],
        [2026, "VIRTUAL", "Búsqueda y uso ético", "Ingeniería", "CAT101", "002", "ABBUEI154-002", "Prof. Ana Gomez", "2026-02-01", "2026-06-16", None, "2026-06-17", "2026-06-18", None],
        [2026, "VIRTUAL", "Búsqueda y uso ético", "Administración", "CAT102", "003", "ABBUEI154-003", "Prof. Carlos Ruiz", "2026-02-01", "2026-06-16", None, "2026-06-17", "2026-06-18", None],
        # Aula ya cerrada (con fecha en certificados)
        [2026, "VIRTUAL", "Búsqueda y uso ético", "Derecho", "CAT103", "001", "ABBUEI100-001", "Prof. Maria Lopez", "2026-01-10", "2026-05-15", None, "2026-05-16", "2026-05-17", "2026-05-18"],
    ]

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    for idx, sname in enumerate(sheet_names):
        if idx == 0:
            ws = wb.active
            ws.title = sname
        else:
            ws = wb.create_sheet(title=sname)

        ws.append(headers)

        if sname == "Tabla_ABBUEI":
            for r in abbuei_rows:
                ws.append(r)
            ws.cell(row=2, column=12).fill = green_fill
            ws.cell(row=2, column=13).fill = green_fill

    wb.save(path)


def create_synthetic_sistematizacion(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sistematización"

    headers = [
        "Modalidad", "Plataforma", "Aula Virtual / Solicitud", "Programa Académico / Dependencia",
        "Fecha Final (DD/MM/AAAA)", "Nombres", "Apellidos", "Tipo de documento de identidad",
        "Número de documento de identidad", "Correo electrónico", "Catálogo", "Clase",
        "Formador líder", "Calificación Módulo 1", "Calificación Módulo 2", "Calificación Módulo 3",
        "Calificación Módulo 4", "Promedio curso completo", "Estado para certificación",
        "Se elaboró certificado SI/NO", "Código del certificado", "Se envió certificado SI/NO", "Fecha de envío",
    ]
    ws.append(headers)

    # Estudiantes ficticios para ABBUEI154
    students = [
        # Clase 001
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Ingeniería", "16/06/2026", "Juan Camilo", "Perez Castro", "CC", "1000000001", "juan.perez1@ejemplo.edu.co", "CAT101", "001", "Prof. Ana Gomez"),
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Ingeniería", "16/06/2026", "Maria Alejandra", "Gomez Toro", "CC", "1000000002", "maria.gomez2@ejemplo.edu.co", "CAT101", "001", "Prof. Ana Gomez"),
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Ingeniería", "16/06/2026", "Carlos Andres", "Restrepo Rios", "CC", "1000000003", "carlos.restrepo3@ejemplo.edu.co", "CAT101", "001", "Prof. Ana Gomez"),
        # Clase 002
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Ingeniería", "16/06/2026", "Laura Marcela", "Velez Jaramillo", "CC", "1000000004", "laura.velez4@ejemplo.edu.co", "CAT101", "002", "Prof. Ana Gomez"),
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Ingeniería", "16/06/2026", "Santiago", "Morales Cano", "CC", "1000000005", "santiago.morales5@ejemplo.edu.co", "CAT101", "002", "Prof. Ana Gomez"),
        # Clase 003
        ("VIRTUAL", "Interactiva", "ABBUEI154", "Administración", "16/06/2026", "Valentina", "Ospina Duque", "CC", "1000000006", "valentina.ospina6@ejemplo.edu.co", "CAT102", "003", "Prof. Carlos Ruiz"),
    ]

    for idx, s in enumerate(students, start=2):
        row = list(s) + [None, None, None, None, None, f'=IF(R{idx}>=3.0, "Aprobado", "No aprobado")', None, None, None, None]
        ws.append(row)

    wb.save(path)


def create_synthetic_notas(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Grades"

    headers = [
        "OrgDefinedId", "Username", "Last Name", "First Name", "Sección",
        "Cuestionario módulo 1: búsqueda", "Buzón módulo 2: evaluación",
        "Cuestionario módulo 3: organización", "Cuestionario módulo 4: ética",
        "Calculated Final Grade",
    ]
    ws.append(headers)

    # Notas con orden alternado y estudiante naranja
    grades_data = [
        # (OrgId, User, Last, First, Sec, M1, M2, M3, M4, Final, is_orange)
        ("ID002", "maria.gomez2", "Gomez Toro", "Maria Alejandra", "001", 4.5, 4.0, 5.0, 4.5, 4.5, False),
        ("ID001", "juan.perez1", "Perez Castro", "Juan Camilo", "001", 5.0, 4.5, 4.8, 5.0, 4.8, False),
        ("ID003", "carlos.restrepo3", "Restrepo Rios", "Carlos Andres", "001", 2.0, 1.5, 2.0, 2.5, 2.0, True),  # Naranja (excluido)
        ("ID005", "santiago.morales5", "Morales Cano", "Santiago", "002", 4.0, 4.2, 4.5, 4.0, 4.2, False),
        ("ID004", "laura.velez4", "Velez Jaramillo", "Laura Marcela", "002", 4.8, 5.0, 4.9, 5.0, 4.9, False),
        ("ID006", "valentina.ospina6", "Ospina Duque", "Valentina", "003", 5.0, 4.8, 5.0, 4.9, 4.9, False),
    ]

    orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")

    for r_idx, g in enumerate(grades_data, start=2):
        row_vals = g[:10]
        ws.append(row_vals)
        if g[10]:  # is_orange
            for c in range(1, 11):
                ws.cell(row=r_idx, column=c).fill = orange_fill

    wb.save(path)


def create_synthetic_certificados(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Certificados"

    headers = [
        "N°", "Ciclo", "Docente COIN", "Nombres", "Apellidos",
        "Documento de identidad", "Correo Electrónico", "Código del certificado",
        "Fecha de envío", "Año de certificación", "Semestre", "Total certificados",
        "Fecha de totalización",
    ]
    ws.append(headers)

    # Registros históricos previos
    prev_rows = [
        [1, "2026-1 / 001", "ABBUEI100 - Prof. Maria Lopez", "Pedro Pablo", "Ramirez Silva", "1000000099", "pedro.ramirez@ejemplo.edu.co", "ABBUEI100", "15-may-2026", 2026, "2026-1", 1, None],
    ]

    for r in prev_rows:
        ws.append(r)

    wb.save(path)


def generate_all_fixtures(output_dir: Path | None = None) -> dict[str, Path]:
    if output_dir is None:
        out = FIXTURES_DIR
    else:
        out = Path(output_dir)

    out.mkdir(parents=True, exist_ok=True)

    paths = {
        "control_aulas": out / "synthetic_control_aulas.xlsx",
        "sistematizacion": out / "synthetic_sistematizacion.xlsx",
        "notas": out / "synthetic_notas.xlsx",
        "certificados": out / "synthetic_certificados.xlsx",
    }

    create_synthetic_control_aulas(paths["control_aulas"])
    create_synthetic_sistematizacion(paths["sistematizacion"])
    create_synthetic_notas(paths["notas"])
    create_synthetic_certificados(paths["certificados"])

    return paths


if __name__ == "__main__":
    generated = generate_all_fixtures()
    print("✅ Fixtures sintéticos generados exitosamente:")
    for k, p in generated.items():
        print(f"   • {k}: {p.name}")
