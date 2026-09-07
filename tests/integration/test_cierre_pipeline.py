"""
Pruebas de Integración de Extremo a Extremo (Pipeline de Cierre de Aulas COIN).

Verifica:
1. Pipeline en Modo Simulación (is_simulation=True): validación en memoria sin generar archivo.
2. Pipeline en Modo Generación (is_simulation=False): generación del Excel auxiliar con las 5 hojas.
3. Respeto estricto del formato oficial en las 5 hojas:
   - RESUMEN
   - SISTEMATIZACION (con notas, promedio y estado calculados)
   - CERTIFICADOS (solo estudiantes con estado 'Aprobó')
   - ACTUALIZACION_CONTROL (propuesta de fecha de cierre)
   - VALIDACIONES_Y_TRAZABILIDAD (auditoría celda a celda)
4. Detección y exclusión correcta de estudiantes con relleno naranja.
5. Robustez absoluta ante filas y columnas ocultas/filtradas en los archivos fuente.
6. Certificación 100% de inmutabilidad de los 4 archivos fuente (SHA-256 intacto).
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from processing.cierre_orchestrator import run_cierre_pipeline
from security.integrity import compute_sha256
from tests.fixtures.generate_fixtures import generate_all_fixtures


class TestCierrePipelineIntegration(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.fixtures = generate_all_fixtures()

    def test_pipeline_simulation_mode(self):
        """Modo 2: Simulación previa en memoria sin escribir archivos."""
        hashes_before = {k: compute_sha256(v).sha256_hash for k, v in self.fixtures.items()}

        result, out_file, intact, msgs = run_cierre_pipeline(
            file_paths=self.fixtures,
            target_aula="ABBUEI154",
            codigo_curso="ABBUEI154",
            is_simulation=True,
        )

        # 1. En simulación NO se genera archivo en disco
        self.assertIsNone(out_file)

        # 2. Integridad 100%
        self.assertTrue(intact)
        for k, v in self.fixtures.items():
            self.assertEqual(compute_sha256(v).sha256_hash, hashes_before[k])

        # 3. Métricas operativas correctas
        self.assertGreater(len(result.clases), 0)
        self.assertEqual(result.total_estudiantes, 6)
        # Hay 1 estudiante con relleno naranja (Carlos Restrepo)
        self.assertEqual(result.total_excluded, 1)
        # Total con notas válidas = 5
        self.assertEqual(result.total_matched, 5)

    def test_pipeline_generation_mode_five_sheets(self):
        """Modo 3: Generación del libro Excel auxiliar de cierre con las 5 hojas estándar."""
        with tempfile.TemporaryDirectory() as tmpdir:
            result, out_file, intact, msgs = run_cierre_pipeline(
                file_paths=self.fixtures,
                target_aula="ABBUEI154",
                codigo_curso="ABBUEI-154",
                is_simulation=False,
                output_dir=tmpdir,
            )

            # 1. Archivo generado existe
            self.assertIsNotNone(out_file)
            self.assertTrue(out_file.exists())
            self.assertTrue(intact)

            # 2. Inspeccionar estructura interna del Excel auxiliar
            wb = openpyxl.load_workbook(out_file, data_only=True)
            expected_sheets = [
                "RESUMEN",
                "SISTEMATIZACION",
                "CERTIFICADOS",
                "ACTUALIZACION_CONTROL",
                "VALIDACIONES_Y_TRAZABILIDAD",
            ]
            for sname in expected_sheets:
                self.assertIn(sname, wb.sheetnames, f"Falta la hoja requerida: {sname}")

            # 3. Hoja 1: RESUMEN
            ws_res = wb["RESUMEN"]
            self.assertIn("INFORME DE CIERRE DE AULA", str(ws_res["A1"].value))
            # Verificar métricas
            res_content = [str(ws_res.cell(r, 1).value or "") for r in range(1, 25)]
            self.assertTrue(any("Total Clases Detectadas" in c for c in res_content))

            # 4. Hoja 2: SISTEMATIZACION
            ws_sist = wb["SISTEMATIZACION"]
            # Debe tener encabezados en fila 1 y filas de datos
            self.assertGreater(ws_sist.max_row, 1)
            # Verificar que las filas de datos no estén ocultas
            for r in range(1, ws_sist.max_row + 1):
                self.assertFalse(ws_sist.row_dimensions[r].hidden)

            # 5. Hoja 3: CERTIFICADOS
            ws_cert = wb["CERTIFICADOS"]
            # Encabezados: N°, Ciclo, Docente COIN, Nombres, Apellidos, ...
            self.assertEqual(ws_cert.cell(1, 1).value, "N°")
            self.assertEqual(ws_cert.cell(1, 2).value, "Ciclo")
            self.assertEqual(ws_cert.cell(1, 3).value, "Docente COIN")

            # Solo estudiantes aprobados deben estar en certificados
            cert_rows = ws_cert.max_row
            # El estudiante naranja (Carlos Restrepo) NO debe estar en certificados
            nombres_cert = [str(ws_cert.cell(r, 4).value or "") for r in range(2, cert_rows + 1)]
            self.assertNotIn("Carlos Andres", nombres_cert)
            # Los estudiantes aprobados SÍ deben estar
            self.assertIn("Juan Camilo", nombres_cert)

            # El consecutivo N° debe ser secuencial
            consecutivos = [ws_cert.cell(r, 1).value for r in range(2, cert_rows + 1)]
            for i, num in enumerate(consecutivos):
                if i > 0:
                    self.assertEqual(num, consecutivos[i - 1] + 1)

            # 6. Hoja 4: ACTUALIZACION_CONTROL
            ws_ctrl = wb["ACTUALIZACION_CONTROL"]
            self.assertGreater(ws_ctrl.max_row, 1)
            self.assertEqual(ws_ctrl.cell(1, 3).value, "Columna Objetivo")
            self.assertEqual(ws_ctrl.cell(2, 3).value, "CERTIFICADOS")

            # 7. Hoja 5: VALIDACIONES_Y_TRAZABILIDAD
            ws_traz = wb["VALIDACIONES_Y_TRAZABILIDAD"]
            self.assertGreater(ws_traz.max_row, 1)
            traz_content = [str(ws_traz.cell(r, 8).value or "") for r in range(2, ws_traz.max_row + 1)]
            # Debe incluir al estudiante excluido
            self.assertTrue(any("EXCLUIDO" in v for v in traz_content))

            wb.close()

    def test_pipeline_with_hidden_and_filtered_rows(self):
        """Verifica que el pipeline procese correctamente archivos con filas y columnas ocultas/filtradas."""
        with tempfile.TemporaryDirectory() as tmpdir:
            # Copiar fixtures a tmpdir y marcar filas/columnas como ocultas
            tmp_path = Path(tmpdir)
            paths = generate_all_fixtures(output_dir=tmp_path)

            # Ocultar fila 2 y columna 8 en Notas (simulando filtro de Excel)
            wb_notas = openpyxl.load_workbook(paths["notas"])
            ws_n = wb_notas.active
            ws_n.row_dimensions[2].hidden = True
            ws_n.column_dimensions["H"].hidden = True
            wb_notas.save(paths["notas"])
            wb_notas.close()

            # Ocultar fila 3 y columna 4 en Sistematización
            wb_sist = openpyxl.load_workbook(paths["sistematizacion"])
            ws_s = wb_sist.active
            ws_s.row_dimensions[3].hidden = True
            ws_s.column_dimensions["D"].hidden = True
            wb_sist.save(paths["sistematizacion"])
            wb_sist.close()

            # Ejecutar pipeline
            result, out_file, intact, msgs = run_cierre_pipeline(
                file_paths=paths,
                target_aula="ABBUEI154",
                codigo_curso="ABBUEI154",
                is_simulation=False,
                output_dir=tmp_path,
            )

            # Debe procesar a todos los 6 estudiantes a pesar de las filas ocultas
            self.assertEqual(result.total_estudiantes, 6)
            self.assertEqual(result.total_matched, 5)
            self.assertEqual(result.total_excluded, 1)

            # En el Excel de trazabilidad debe constar la advertencia informativa de fila oculta
            wb_out = openpyxl.load_workbook(out_file, data_only=True)
            ws_t = wb_out["VALIDACIONES_Y_TRAZABILIDAD"]
            adv_texts = [str(ws_t.cell(r, 9).value or "") for r in range(2, ws_t.max_row + 1)]
            self.assertTrue(
                any("Fila oculta" in a for a in adv_texts),
                "No se registró la trazabilidad de fila oculta",
            )
            wb_out.close()


if __name__ == "__main__":
    unittest.main()
