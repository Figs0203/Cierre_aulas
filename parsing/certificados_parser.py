"""
Parser especializado para el archivo de Certificados (Control_codigos_certificados_*.xlsx).

Inspecciona la hoja 'Códigos' para determinar:
- El último número correlativo (N°) utilizado, para sugerir la numeración correcta.
- La cantidad de certificados previamente registrados para el aula a cerrar.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Tuple

import openpyxl


def parse_certificados_consecutivo(
    file_path: Path | str,
    target_aula: Optional[str] = None,
    sheet_name: str = "Códigos",
) -> Tuple[int, int]:
    """Extrae el último consecutivo N° y el número de certificados existentes del aula.

    Args:
        file_path: Ruta al archivo de certificados.
        target_aula: Código del aula a comprobar (opcional).
        sheet_name: Nombre de la hoja de certificados (por defecto 'Códigos').

    Returns:
        Tupla de:
          - max_consecutivo: Último número correlativo entero encontrado en Col A (0 si no hay).
          - prev_cert_count: Cantidad de registros previos que ya tienen este target_aula en Col H.
    """
    path = Path(file_path)
    clean_target = target_aula.strip().upper() if target_aula else None

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif "Certificados" in wb.sheetnames:
        ws = wb["Certificados"]
    else:
        ws = wb.sheetnames[0]
        ws = wb[ws]

    max_n = 0
    prev_count = 0
    max_row = min(ws.max_row or 1, 30000)

    for r in range(2, max_row + 1):
        # Col 1 (A): N°
        val_n = ws.cell(r, 1).value
        if val_n is not None:
            try:
                n_int = int(float(str(val_n).strip()))
                if n_int > max_n:
                    max_n = n_int
            except (ValueError, TypeError):
                pass

        # Col 8 (H): Código del certificado
        if clean_target:
            val_cod = ws.cell(r, 8).value
            if val_cod and clean_target in str(val_cod).strip().upper():
                prev_count += 1

    wb.close()
    return max_n, prev_count
