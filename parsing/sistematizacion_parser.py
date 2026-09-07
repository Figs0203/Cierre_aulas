"""
Parser especializado para el archivo de Sistematización de Cursos COIN.

Extrae información de estudiantes y calificaciones existentes desde la
hoja 'Formato' de Sistematizacion_Cursos_COIN_2026.xlsx, filtrando por el
código del aula virtual (ej. 'ABBUEI205').

Mapeo confirmado por Informe de Descubrimiento de Fase 0 (2026-09-07):
  Hoja: 'Formato' (55 columnas)
  Col  4 ( D): Tipo de grupo ('Pregrado' / 'Posgrado')
  Col  5 ( E): Aula Virtual / Solicitud
  Col 25 ( Y): Nombres
  Col 26 ( Z): Apellidos
  Col 27 (AA): Tipo de documento de identidad
  Col 28 (AB): Número de documento de identidad
  Col 29 (AC): Correo electrónico
  Col 35 (AI): Catálogo
  Col 36 (AJ): Clase
  Col 37 (AK): Formador líder
  Col 38 (AL): Formador acompañante
  Col 45 (AS): Calificación Módulo 1
  Col 46 (AT): Calificación Módulo 2
  Col 47 (AU): Calificación Módulo 3
  Col 48 (AV): Calificación Módulo 4
  Col 49 (AW): Calificación Módulo 5
  Col 50 (AX): Promedio curso completo
  Col 51 (AY): Estado para certificación
  Col 52 (AZ): Se elaboró certificado SI/NO
  Col 53 (BA): Código del certificado
  Col 54 (BB): Se envió certificado SI/NO
  Col 55 (BC): Fecha de envío (DD/MM/AAAA)
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from core.models import EstudianteSistematizacion
from core.traceability import TracedValue
from normalization.text import normalize_document, normalize_email


def parse_sistematizacion_aula(
    file_path: Path | str,
    target_aula: str,
    sheet_name: str = "Formato",
) -> Tuple[Dict[str, List[EstudianteSistematizacion]], List[str]]:
    """Extrae los estudiantes de un aula virtual desde la hoja de Sistematización.

    Args:
        file_path: Ruta al archivo Excel de Sistematización.
        target_aula: Código del aula a filtrar (ej. 'ABBUEI205').
        sheet_name: Nombre de la hoja operativa (por defecto 'Formato').

    Returns:
        Tupla de:
          - Dict {clase_id: list[EstudianteSistematizacion]}
          - Lista de encabezados originales de la hoja (para reconstrucción exacta)
    """
    path = Path(file_path)
    clean_target = target_aula.strip().upper()
    file_name = path.name

    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)

    # Si la hoja 'Formato' no existe, buscar una hoja alternativa adecuada
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    elif "Sistematización" in wb.sheetnames:
        sheet_name = "Sistematización"
        ws = wb[sheet_name]
    else:
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]

    # 1. Leer encabezados de la fila 1
    max_col = ws.max_column or 55
    raw_headers: List[str] = []
    header_map: Dict[str, int] = {}

    for c in range(1, max_col + 1):
        v = ws.cell(1, c).value
        h_str = str(v).strip() if v is not None else f"Col_{c}"
        raw_headers.append(h_str)
        # Normalizar para búsqueda flexible si las columnas se mueven
        h_norm = h_str.lower().replace("\n", " ")
        header_map[h_norm] = c

    # 2. Localizar índices de columnas clave
    # Si coinciden con el estándar de 55 columnas de la Fase 0, usar posiciones exactas
    col_aula = 5
    col_tipo_grupo = 4
    col_nombres = 25
    col_apellidos = 26
    col_tipo_doc = 27
    col_doc = 28
    col_correo = 29
    col_catalogo = 35
    col_clase = 36
    col_formador = 37
    col_formador_acomp = 38
    col_m1 = 45
    col_m2 = 46
    col_m3 = 47
    col_m4 = 48
    col_m5 = 49
    col_promedio = 50
    col_estado = 51
    col_elaboro = 52
    col_cod_cert = 53
    col_envio = 54
    col_fecha_envio = 55

    # Verificación dinámica en caso de formato alternativo o columnas movidas
    for h_name, c_idx in header_map.items():
        if "aula virtual" in h_name or "solicitud" in h_name:
            col_aula = c_idx
        elif "tipo de grupo" in h_name:
            col_tipo_grupo = c_idx
        elif "nombres" == h_name:
            col_nombres = c_idx
        elif "apellidos" == h_name:
            col_apellidos = c_idx
        elif "número de documento" in h_name or "numero de documento" in h_name:
            col_doc = c_idx
        elif "correo" in h_name:
            col_correo = c_idx
        elif "catálogo" in h_name or "catalogo" in h_name:
            col_catalogo = c_idx
        elif "clase" == h_name:
            col_clase = c_idx
        elif "formador líder" in h_name or "formador lider" in h_name:
            col_formador = c_idx
        elif "formador acompañante" in h_name or "formador acompanante" in h_name:
            col_formador_acomp = c_idx
        elif "calificación módulo 1" in h_name or "calificacion modulo 1" in h_name or "módulo 1" in h_name or "modulo 1" in h_name:
            col_m1 = c_idx
        elif "calificación módulo 2" in h_name or "calificacion modulo 2" in h_name or "módulo 2" in h_name or "modulo 2" in h_name:
            col_m2 = c_idx
        elif "calificación módulo 3" in h_name or "calificacion modulo 3" in h_name or "módulo 3" in h_name or "modulo 3" in h_name:
            col_m3 = c_idx
        elif "calificación módulo 4" in h_name or "calificacion modulo 4" in h_name or "módulo 4" in h_name or "modulo 4" in h_name:
            col_m4 = c_idx
        elif "calificación módulo 5" in h_name or "calificacion modulo 5" in h_name or "módulo 5" in h_name or "modulo 5" in h_name:
            col_m5 = c_idx
        elif "promedio" in h_name:
            col_promedio = c_idx
        elif "estado para certificación" in h_name or "estado certificación" in h_name or "estado certificacion" in h_name:
            col_estado = c_idx
        elif "se elaboró" in h_name or "se elaboro" in h_name:
            col_elaboro = c_idx
        elif "código del certificado" in h_name or "codigo del certificado" in h_name or "código certificado" in h_name or "codigo certificado" in h_name:
            col_cod_cert = c_idx
        elif "se envió certificado" in h_name or "se envio certificado" in h_name:
            col_envio = c_idx
        elif "fecha de envío" in h_name or "fecha de envio" in h_name:
            col_fecha_envio = c_idx

    dynamic_column_indices = {
        "aula": col_aula,
        "tipo_grupo": col_tipo_grupo,
        "nombres": col_nombres,
        "apellidos": col_apellidos,
        "tipo_doc": col_tipo_doc,
        "doc": col_doc,
        "correo": col_correo,
        "catalogo": col_catalogo,
        "clase": col_clase,
        "formador": col_formador,
        "formador_acomp": col_formador_acomp,
        "modulo_1": col_m1,
        "modulo_2": col_m2,
        "modulo_3": col_m3,
        "modulo_4": col_m4,
        "modulo_5": col_m5,
        "promedio": col_promedio,
        "estado": col_estado,
        "elaboro": col_elaboro,
        "codigo_cert": col_cod_cert,
        "envio": col_envio,
        "fecha_envio": col_fecha_envio,
    }

    # 3. Iterar y extraer filas que coincidan con el aula objetivo
    estudiantes_por_clase: Dict[str, List[EstudianteSistematizacion]] = {}
    max_row = ws.max_row or 1

    for r in range(2, max_row + 1):
        cell_aula = ws.cell(r, col_aula).value
        if not cell_aula:
            continue
        if str(cell_aula).strip().upper() != clean_target:
            continue

        # Detectar si la fila estaba oculta o filtrada en Excel (pero procesarla sin omisiones)
        is_hidden = bool(ws.row_dimensions[r].hidden)

        # Extraer toda la fila bruta para reproducibilidad 100%
        raw_row_data = {c: ws.cell(r, c).value for c in range(1, max_col + 1)}

        # Clase
        clase_val = ws.cell(r, col_clase).value
        clase_str = str(clase_val).strip() if clase_val is not None else "001"
        # Si la clase tiene formato flotante como 6872.0, convertir a entero
        if isinstance(clase_val, float) and clase_val.is_integer():
            clase_str = str(int(clase_val))

        clase_traced = TracedValue(
            value=clase_str,
            source_file=file_name,
            source_sheet=sheet_name,
            source_column="Clase",
            source_col_idx=col_clase,
            source_row=r,
        )

        def make_traced(val: Any, col_name: str, c_idx: int) -> TracedValue:
            return TracedValue(
                value=val,
                source_file=file_name,
                source_sheet=sheet_name,
                source_column=col_name,
                source_col_idx=c_idx,
                source_row=r,
            )

        nombres_v = ws.cell(r, col_nombres).value
        apellidos_v = ws.cell(r, col_apellidos).value
        tipo_doc_v = ws.cell(r, col_tipo_doc).value
        doc_v = ws.cell(r, col_doc).value
        correo_v = ws.cell(r, col_correo).value
        tipo_grupo_v = ws.cell(r, col_tipo_grupo).value
        catalogo_v = ws.cell(r, col_catalogo).value
        formador_v = ws.cell(r, col_formador).value
        formador_acomp_v = ws.cell(r, col_formador_acomp).value
        estado_v = ws.cell(r, col_estado).value

        # Calificaciones existentes
        califs_existentes = {}
        for m_num, m_col in [("modulo_1", col_m1), ("modulo_2", col_m2), ("modulo_3", col_m3), ("modulo_4", col_m4), ("modulo_5", col_m5)]:
            if m_col <= max_col:
                cv = ws.cell(r, m_col).value
                if cv is not None:
                    califs_existentes[m_num] = make_traced(cv, f"Calificación Módulo {m_num[-1]}", m_col)

        promedio_v = ws.cell(r, col_promedio).value if col_promedio <= max_col else None

        est = EstudianteSistematizacion(
            row_number=r,
            is_hidden_row=is_hidden,
            column_indices=dynamic_column_indices,
            nombres=make_traced(nombres_v, "Nombres", col_nombres) if nombres_v is not None else None,
            apellidos=make_traced(apellidos_v, "Apellidos", col_apellidos) if apellidos_v is not None else None,
            tipo_documento=make_traced(tipo_doc_v, "Tipo de documento de identidad", col_tipo_doc) if tipo_doc_v is not None else None,
            documento=make_traced(doc_v, "Número de documento de identidad", col_doc) if doc_v is not None else None,
            correo=make_traced(correo_v, "Correo electrónico", col_correo) if correo_v is not None else None,
            clase=clase_traced,
            catalogo=make_traced(catalogo_v, "Catálogo", col_catalogo) if catalogo_v is not None else None,
            tipo_grupo=make_traced(tipo_grupo_v, "Tipo de grupo", col_tipo_grupo) if tipo_grupo_v is not None else None,
            formador_lider=make_traced(formador_v, "Formador líder", col_formador) if formador_v is not None else None,
            formador_acompanante=make_traced(formador_acomp_v, "Formador acompañante", col_formador_acomp) if formador_acomp_v is not None else None,
            calificaciones_existentes=califs_existentes,
            promedio_existente=make_traced(promedio_v, "Promedio curso completo", col_promedio) if promedio_v is not None else None,
            estado_certificacion_raw=make_traced(estado_v, "Estado para certificación", col_estado) if estado_v is not None else None,
            formula_estado_certificacion=str(estado_v) if isinstance(estado_v, str) and estado_v.startswith("=") else None,
            raw_row_data=raw_row_data,
        )

        if clase_str not in estudiantes_por_clase:
            estudiantes_por_clase[clase_str] = []
        estudiantes_por_clase[clase_str].append(est)

    wb.close()
    return estudiantes_por_clase, raw_headers
